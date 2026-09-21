#!/usr/bin/env python3
"""Step 3d Phase 1 follow-up: the 6th "Define KPI Framework" rename.

Discovered after the 8-rename Phase 1 execution (PR #84): CM-1-3-3-5-7
under Marketing Communications Strategy still carried the generic
"Define KPI Framework" prefLabel, with "Define Marketing Communications
Measurement Framework" already queued as its altLabel.

Hamid's decision (2026-09-21): rename to
"Define Marketing Communications Measurement Framework" so it matches
the five siblings — the generic label must not survive as the one
remaining semantic collision.

Treatment matches the five sibling KPI rows exactly:
  - prefLabel <- "Define Marketing Communications Measurement Framework"
  - the old generic label is NOT retained as an ontology altLabel
    (it would recreate the collision; scoped historical alias
    "Define KPI Framework (Marketing Communications)" is recorded in
    the identity/migration map instead)
  - execution record appended to the row's terminology_notes

One-shot. Asserts expected pre-state, so a stale workbook fails loudly.
Run from the repo root. Adds the row to the naming queue's
critical-collision section (it was never queued) marked EXECUTED.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[4]
WB = ROOT / "business_architecture/ontology/build/output/step3c-definition-authoring-workbook.xlsx"
SHEET = "Review & authoring"
IDENTITY_MAP = ROOT / "business_architecture/ontology/build/output/step2-identity-map.json"
QUEUE = ROOT / "business_architecture/ontology/step3c-naming-pass-queue.md"

SLUG = "CM-1-3-3-5-7"
OLD_NAME = "Define KPI Framework"
NEW_NAME = "Define Marketing Communications Measurement Framework"
SCOPED_ALIAS = "Define KPI Framework (Marketing Communications)"

EXEC_NOTE = (
    "Step 3d naming pass (2026-09-21): renamed from 'Define KPI Framework' to "
    "'Define Marketing Communications Measurement Framework' — sixth and final "
    "'Define KPI Framework' collision resolved (Hamid's follow-up decision, "
    "discovered after the 8-rename Phase 1 execution). Prior label retired from "
    "ontology aliases (a shared generic altLabel would recreate the collision); "
    f"scoped historical alias '{SCOPED_ALIAS}' recorded in the identity/migration map."
)


def main() -> int:
    # 1. workbook
    wb = load_workbook(WB)
    ws = wb[SHEET]
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    row = None
    for r in ws.iter_rows(min_row=2):
        if r[idx["slug"]].value == SLUG:
            row = r
            break
    assert row is not None, f"{SLUG} not found in workbook"
    old_name = row[idx["name"]].value
    assert old_name == OLD_NAME, f"{SLUG}: name {old_name!r} != expected {OLD_NAME!r}"
    old_alt = row[idx["alt_labels"]].value or ""
    assert old_alt == NEW_NAME, f"{SLUG}: alt {old_alt!r} != expected {NEW_NAME!r}"
    row[idx["name"]].value = NEW_NAME
    row[idx["alt_labels"]].value = None  # generic retired, like the five siblings
    cur = row[idx["terminology_notes"]].value or ""
    row[idx["terminology_notes"]].value = (cur + " " + EXEC_NOTE).strip() if cur else EXEC_NOTE
    wb.save(WB)
    print(f"workbook: renamed {SLUG}: {OLD_NAME!r} -> {NEW_NAME!r}")

    # 2. identity map
    entries = json.loads(IDENTITY_MAP.read_text(encoding="utf-8"))
    entry = next(e for e in entries if e["slug"] == SLUG)
    assert entry["name"] == OLD_NAME, f"identity map name {entry['name']!r} != {OLD_NAME!r}"
    entry["name"] = NEW_NAME
    entry["prior_name"] = OLD_NAME
    entry["scoped_historical_alias"] = SCOPED_ALIAS
    entry["name_change_note"] = (
        "Step 3d naming pass (2026-09-21): renamed 'Define KPI Framework' -> "
        f"'{NEW_NAME}' — sixth and final 'Define KPI Framework' collision "
        "resolved (approved by Hamid). Prior label retired from ontology "
        f"aliases; scoped historical alias '{SCOPED_ALIAS}' recorded here."
    )
    IDENTITY_MAP.write_text(json.dumps(entries, indent=1, ensure_ascii=False) + "\n",
                            encoding="utf-8")
    print(f"identity map: updated {SLUG} + migration metadata")

    # 3. naming queue: add the late-discovered row as the 9th critical entry
    text = QUEUE.read_text(encoding="utf-8")
    assert "| `CM-1-3-3-6-5` | Define KPI Framework |" in text, "queue sibling row missing"
    new_row = (
        "| `CM-1-3-3-5-7` | Define KPI Framework | Define Marketing Communications "
        "Measurement Framework **[priority]** | Template row (batch 23 treatment); "
        "preferred label 'Define Marketing Communications Measurement Framework' queued "
        "prominently; full measurement-framework routing statement in scope note per review. "
        "Sixth 'Define KPI Framework' collision, discovered after the 8-rename Phase 1 "
        "execution. **EXECUTED 2026-09-21** (Step 3d Phase 1 follow-up: renamed to "
        "'Define Marketing Communications Measurement Framework'; see naming PR). |"
    )
    anchor = "| `CM-1-3-3-6-5` | Define KPI Framework | Define Operating Model Measurement Framework **[priority]** |"
    lines = text.split("\n")
    for i, line in enumerate(lines):
        if line.startswith(anchor):
            lines.insert(i + 1, new_row)
            break
    else:
        raise AssertionError("anchor row not found")
    text = "\n".join(lines)
    old_hdr = ("**Execution status 2026-09-21:** all 8 critical-collision entries executed in Step 3d Phase 1\n"
               "(Hamid approval 2026-09-21). The remaining 84 queue entries are still queued for the consolidated naming pass.")
    new_hdr = ("**Execution status 2026-09-21:** all 9 critical-collision entries executed in Step 3d Phase 1\n"
               "(Hamid approval 2026-09-21, including the 6th 'Define KPI Framework' row discovered after\n"
               "the initial 8-rename execution). The remaining 84 queue entries are still queued for the consolidated naming pass.")
    assert old_hdr in text, "queue header not in expected form"
    text = text.replace(old_hdr, new_hdr)
    old_sec = "## Critical semantic collision (8)"
    assert old_sec in text
    text = text.replace(old_sec, "## Critical semantic collision (9)")
    QUEUE.write_text(text, encoding="utf-8")
    print("naming queue: added 9th critical-collision entry, marked executed")

    return 0


if __name__ == "__main__":
    sys.exit(main())
