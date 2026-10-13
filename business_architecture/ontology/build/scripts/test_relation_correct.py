#!/usr/bin/env python3
"""Fixture tests for relation-correct.py."""
from __future__ import annotations

import csv
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

SCRIPT = Path(__file__).resolve().parent / "relation-correct.py"
HEADERS = ["slug", "name", "related_concepts", "status"]
COLUMNS = ["row_id", "source_slug", "old_verb", "old_target", "new_verb", "new_target"]
CELL = "enables: Trade Capture (CM-1-2-1-3) | governed-by: Market Risk Management (CM-1-2-2-3) | informs: Plan (CM-9)"


def _fixture(cell: str, edits: list[list[str]], status: str = "approved") -> tuple[Path, Path]:
    tmp = Path(tempfile.mkdtemp())
    wb = Workbook()
    ws = wb.active
    ws.title = "Review & authoring"
    ws.append(HEADERS)
    ws.append(["A-1", "Source", cell, status])
    wb.save(tmp / "wb.xlsx")
    with open(tmp / "c.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(COLUMNS)
        w.writerows(edits)
    return tmp / "wb.xlsx", tmp / "c.csv"


def _run(wb: Path, c: Path, *extra: str) -> subprocess.CompletedProcess:
    return subprocess.run([sys.executable, str(SCRIPT), "--workbook", str(wb), "--corrections", str(c), *extra],
                          capture_output=True, text=True)


def _cell(wb: Path) -> str:
    return next(load_workbook(wb)["Review & authoring"].iter_rows(min_row=2, values_only=True))[2]


class RelationCorrect(unittest.TestCase):
    def test_rewrites_one_segment_in_place_and_keeps_the_others(self):
        wb, c = _fixture(CELL, [["REL-1", "A-1", "enables", "Trade Capture (CM-1-2-1-3)",
                                 "informs", "Trade Capture (CM-1-2-1-3)"]])
        self.assertEqual(_run(wb, c).returncode, 0)
        self.assertEqual(_cell(wb), CELL.replace("enables: Trade Capture", "informs: Trade Capture"))
        self.assertEqual(_cell(wb).count("|"), CELL.count("|"))

    def test_retarget_changes_only_the_target(self):
        wb, c = _fixture(CELL, [["REL-2", "A-1", "governed-by", "Market Risk Management (CM-1-2-2-3)",
                                 "governed-by", "Establish DOA (CM-1-2-2-3-2)"]])
        self.assertEqual(_run(wb, c).returncode, 0)
        self.assertIn(" governed-by: Establish DOA (CM-1-2-2-3-2) ", _cell(wb))

    def test_rerun_changes_nothing(self):
        edit = [["REL-1", "A-1", "enables", "Trade Capture (CM-1-2-1-3)", "informs", "Trade Capture (CM-1-2-1-3)"]]
        wb, c = _fixture(CELL, edit)
        _run(wb, c)
        after = _cell(wb)
        result = _run(wb, c)
        self.assertIn("1 already corrected", result.stdout)
        self.assertEqual(_cell(wb), after)

    def test_unchanged_row_is_ignored(self):
        wb, c = _fixture(CELL, [["REL-1", "A-1", "enables", "Trade Capture (CM-1-2-1-3)",
                                 "enables", "Trade Capture (CM-1-2-1-3)"]])
        self.assertIn("0 changed", _run(wb, c).stdout)
        self.assertEqual(_cell(wb), CELL)

    def test_refuses_missing_or_duplicate_segment_before_any_write(self):
        dup = CELL + " | enables: Trade Capture (CM-1-2-1-3)"
        for cell, old in ((CELL, "Nowhere (CM-0)"), (dup, "Trade Capture (CM-1-2-1-3)")):
            wb, c = _fixture(cell, [["REL-1", "A-1", "enables", old, "informs", old]])
            result = _run(wb, c)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("refused", result.stderr)
            self.assertEqual(_cell(wb), cell)

    def test_refuses_duplicate_or_mixed_corrected_segments_before_any_write(self):
        good = ["REL-2", "A-1", "informs", "Plan (CM-9)", "enables", "Plan (CM-9)"]
        dup = CELL.replace("enables: Trade Capture", "informs: Trade Capture") + " | informs: Trade Capture (CM-1-2-1-3)"
        both = CELL + " | informs: Trade Capture (CM-1-2-1-3)"
        for cell in (dup, both):
            wb, c = _fixture(cell, [["REL-1", "A-1", "enables", "Trade Capture (CM-1-2-1-3)",
                                     "informs", "Trade Capture (CM-1-2-1-3)"], good])
            before = wb.read_bytes()
            result = _run(wb, c)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("already appears", result.stderr)
            self.assertEqual(wb.read_bytes(), before)

    def test_refuses_unknown_verb_and_unapproved_row(self):
        wb, c = _fixture(CELL, [["REL-1", "A-1", "enables", "Trade Capture (CM-1-2-1-3)",
                                 "feeds", "Trade Capture (CM-1-2-1-3)"]])
        self.assertIn("not in the direction table", _run(wb, c).stderr)
        wb, c = _fixture(CELL, [["REL-1", "A-1", "enables", "Trade Capture (CM-1-2-1-3)",
                                 "informs", "Trade Capture (CM-1-2-1-3)"]], status="pending")
        self.assertIn("not 'approved'", _run(wb, c).stderr)
        self.assertEqual(_cell(wb), CELL)

    def test_accepts_the_mirror_verbs(self):
        wb, c = _fixture(CELL, [["REL-1", "A-1", "enables", "Trade Capture (CM-1-2-1-3)",
                                 "provides-input-to", "Trade Capture (CM-1-2-1-3)"],
                                ["REL-2", "A-1", "governed-by", "Market Risk Management (CM-1-2-2-3)",
                                 "governs", "Market Risk Management (CM-1-2-2-3)"]])
        self.assertEqual(_run(wb, c).returncode, 0)
        self.assertIn("provides-input-to: Trade Capture", _cell(wb))
        self.assertIn("governs: Market Risk Management", _cell(wb))

    def test_check_writes_nothing(self):
        wb, c = _fixture(CELL, [["REL-1", "A-1", "enables", "Trade Capture (CM-1-2-1-3)",
                                 "informs", "Trade Capture (CM-1-2-1-3)"]])
        self.assertIn("1 changed", _run(wb, c, "--check").stdout)
        self.assertEqual(_cell(wb), CELL)


if __name__ == "__main__":
    unittest.main()
