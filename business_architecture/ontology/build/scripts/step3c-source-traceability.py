#!/usr/bin/env python3
"""Generate the source-to-process traceability appendix for the business
process architecture playbook.

Reads the Step 3c workbook (Reference register + Review & authoring) and
emits, for every registered source, the process areas and rows that cite it,
alongside the source's permitted-use limits. Regenerate after every batch:

    python3 business_architecture/ontology/build/scripts/step3c-source-traceability.py

Output: business_architecture/ontology/build/output/source-to-process-traceability.md
Read-only over the workbook; deterministic; safe to re-run.
"""
import datetime
import os
from collections import defaultdict

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.normpath(os.path.join(HERE, ".."))
WB = os.path.join(ROOT, "output", "step3c-definition-authoring-workbook.xlsx")
OUT = os.path.join(ROOT, "output", "source-to-process-traceability.md")


def area_of(breadcrumb, slug):
    """A readable process-area label: the L2/L3 stretch of the breadcrumb."""
    if breadcrumb:
        parts = [p.strip() for p in str(breadcrumb).split(">")]
        if len(parts) >= 4:
            return " > ".join(parts[1:4])
        if len(parts) >= 2:
            return " > ".join(parts[1:])
    return slug or "(unknown)"


def main():
    wb = load_workbook(WB, read_only=True)

    reg = wb["Reference register"]
    reg_rows = list(reg.iter_rows(values_only=True))
    hdr_i = next(i for i, r in enumerate(reg_rows) if r and r[0] == "source_id")
    sources = {}
    order = []
    for r in reg_rows[hdr_i + 1:]:
        if r and r[0]:
            sources[r[0]] = {
                "organization": r[1] or "",
                "title": r[2] or "",
                "version_or_date": r[3] or "",
                "url": r[4] or "",
                "authority_type": r[5] or "",
                "permitted_use": r[6] or "",
                "status": r[8] or "",
            }
            order.append(r[0])

    ws = wb["Review & authoring"]
    it = ws.iter_rows(values_only=True)
    hdr = {v: i for i, v in enumerate(next(it))}
    cites = defaultdict(lambda: defaultdict(list))  # source -> area -> [(slug, name)]
    unknown = defaultdict(list)
    n_rows_citing = 0
    for r in it:
        refs = r[hdr["reference_sources"]]
        if not refs:
            continue
        slug, name = r[hdr["slug"]], r[hdr["name"]]
        area = area_of(r[hdr["breadcrumb"]], slug)
        n_rows_citing += 1
        for sid in [s.strip().rstrip(".") for s in str(refs).replace("|", ";").split(";")]:
            if not sid:
                continue
            sid = sid.split(" ")[0]
            if sid in sources:
                cites[sid][area].append((slug, name))
            else:
                unknown[sid].append(slug)

    lines = [
        "# Source-to-Process Traceability",
        "",
        f"**Generated:** {datetime.date.today().isoformat()} by "
        "`build/scripts/step3c-source-traceability.py` — do not edit by hand; regenerate instead.",
        f"**Workbook:** `{os.path.relpath(WB, ROOT)}` · {len(sources)} registered sources · "
        f"{n_rows_citing} rows carry citations.",
        "",
        "Each registered source below lists the process areas and rows whose definitions it "
        "informs, with the source's permitted-use limits. Sources are evidence, not authority: "
        "permitted-use text records what each source may and may not determine.",
        "",
    ]
    cited = [s for s in order if s in cites]
    uncited = [s for s in order if s not in cites]

    for sid in cited:
        meta = sources[sid]
        lines.append(f"## {sid}")
        lines.append("")
        lines.append(f"**{meta['organization']}** — {meta['title']}")
        lines.append("")
        lines.append(f"- **Reference:** {meta['url']}")
        lines.append(f"- **Version/access:** {meta['version_or_date']}")
        lines.append(f"- **Authority type:** {meta['authority_type']} · **Status:** {meta['status']}")
        lines.append(f"- **Permitted use:** {meta['permitted_use']}")
        lines.append("")
        lines.append("**Informs:**")
        lines.append("")
        for area in sorted(cites[sid]):
            rows = cites[sid][area]
            lines.append(f"- *{area}*")
            for slug, name in rows:
                lines.append(f"  - `{slug}` — {name}")
        lines.append("")

    if uncited:
        lines.append("## Registered sources not currently cited from any row")
        lines.append("")
        lines.append("Retained in the register (registered for earlier steps, or cited from "
                     "non-workbook artifacts):")
        lines.append("")
        for sid in uncited:
            meta = sources[sid]
            lines.append(f"- **{sid}** — {meta['organization']}: {meta['title']}")
        lines.append("")

    if unknown:
        lines.append("## ⚠ Citations with no register entry (fix these)")
        lines.append("")
        for sid, slugs in sorted(unknown.items()):
            lines.append(f"- `{sid}` cited from: {', '.join(slugs)}")
        lines.append("")

    with open(OUT, "w") as f:
        f.write("\n".join(lines))
    print(f"wrote {OUT}: {len(cited)} cited sources, {len(uncited)} uncited, "
          f"{sum(len(v) for v in unknown.values())} unknown citations")


if __name__ == "__main__":
    main()
