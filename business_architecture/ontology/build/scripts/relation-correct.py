#!/usr/bin/env python3
"""Apply approved relationship corrections to the Step 3c definition workbook.

Phase 1 source corrections (EPM-DEC-001-0025) change a relationship's verb or
target in the source, never in the promotion script (rule 11: emission never
changes source meaning). Each relationship is one `verb: target` segment of a
process's `related_concepts` cell on the "Review & authoring" sheet, and Step 4
numbers its REL rows by segment position. So a correction rewrites one segment
in place and never adds, removes or moves one: every REL ID stays the same.

The input is a corrections CSV with columns row_id, source_slug, old_verb,
old_target, new_verb, new_target. A row whose new values equal its old values
changes nothing.

Rules:
- The old segment must appear exactly once in the source's cell, verb and
  target text matching; otherwise the run stops before any write.
- The new verb must be a raw verb the Step 4 direction table knows.
- The new target must be printable text without '|', '"' or '\\', and every
  corrected cell keeps its number of segments.
- Only `approved` rows are edited, because only they reach the taxonomy.
- A cell already holding the new values once, and the old values nowhere, is
  skipped, so a re-run changes nothing. Any other mix of old and new
  segments stops the run.

Usage:
  relation-correct.py --corrections <csv> [--workbook <xlsx>] [--check]
--check reports what would change and writes nothing.
"""
import argparse
import csv
import sys
from pathlib import Path

WB = Path(__file__).resolve().parents[1] / "output" / "step3c-definition-authoring-workbook.xlsx"
SHEET = "Review & authoring"
# Raw verbs in verb-predicate-mapping-v2.md's canonical direction table.
VERBS = {"enables", "informs", "informed-by", "uses-input", "governed-by", "requires", "constrained-by",
         "constrains", "triggers", "assures", "precedes", "follows",
         "provides-input-to", "governs"}  # the two mirror verbs: EPM-DEC-001-0027
COLUMNS = ["row_id", "source_slug", "old_verb", "old_target", "new_verb", "new_target"]


def segments(cell):
    """The cell split on '|', keeping each segment's own spacing."""
    return (cell or "").split("|")


def parse(segment):
    if ":" not in segment:
        return None
    verb, target = segment.split(":", 1)
    return verb.strip(), target.strip()


def corrections(path):
    with open(path, newline="") as f:
        rows = list(csv.DictReader(f))
    if rows and not set(COLUMNS) <= rows[0].keys():
        sys.exit(f"{path}: needs columns {', '.join(COLUMNS)}")
    return [{k: " ".join((r[k] or "").split()) for k in COLUMNS} for r in rows]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--corrections", required=True)
    ap.add_argument("--workbook", default=str(WB))
    ap.add_argument("--check", action="store_true")
    args = ap.parse_args()

    from openpyxl import load_workbook

    wb = load_workbook(args.workbook)
    ws = wb[SHEET]
    idx = {c.value: i for i, c in enumerate(ws[1])}
    by_slug = {r[idx["slug"]].value: r for r in ws.iter_rows(min_row=2) if r[idx["slug"]].value}

    edits = [e for e in corrections(args.corrections)
             if (e["old_verb"], e["old_target"]) != (e["new_verb"], e["new_target"])]
    errors, plan = [], []
    for e in edits:
        rid, slug = e["row_id"], e["source_slug"]
        if e["new_verb"] not in VERBS:
            errors.append(f"{rid}: new verb {e['new_verb']!r} is not in the direction table")
        if "|" in e["new_target"]:
            errors.append(f"{rid}: new target contains '|', which would add a segment")
        # No workbook target holds a quote or backslash, and corrections keep it that way:
        # a target is a plain process name. The classifier reads escapes since #240, so this
        # is policy and defence in depth, not a workaround for a misread.
        if set(e["new_target"]) & {'"', "\\"}:
            errors.append(f"{rid}: new target contains a quote or backslash; a target is a plain process name")
        # Non-printable characters (controls, U+FFFE, U+FFFF) cannot be saved as worksheet XML.
        if not e["new_target"].isprintable():
            errors.append(f"{rid}: new target contains a non-printable character")
        if slug not in by_slug:
            errors.append(f"{rid}: {slug} not in the workbook")
            continue
        if (by_slug[slug][idx["status"]].value or "").strip() != "approved":
            errors.append(f"{rid}: {slug} status is {by_slug[slug][idx['status']].value!r}, not 'approved'")
            continue
        parts = segments(by_slug[slug][idx["related_concepts"]].value)
        new = [i for i, p in enumerate(parts) if parse(p) == (e["new_verb"], e["new_target"])]
        old = [i for i, p in enumerate(parts) if parse(p) == (e["old_verb"], e["old_target"])]
        if len(new) == 1 and not old:
            plan.append((e, None))
        elif new:
            errors.append(f"{rid}: '{e['new_verb']}: {e['new_target']}' already appears {len(new)} times in {slug}'s cell")
        elif len(old) != 1:
            errors.append(f"{rid}: '{e['old_verb']}: {e['old_target']}' appears {len(old)} times in {slug}'s cell")
        else:
            plan.append((e, old[0]))
    if errors:
        sys.exit("refused, nothing written:\n  " + "\n  ".join(errors))

    # Apply every plan to copies of the cells, then check the combined result, so two
    # corrections cannot together leave a duplicate that a re-run would refuse.
    final = {e["source_slug"]: segments(by_slug[e["source_slug"]][idx["related_concepts"]].value) for e, _ in plan}
    changed = skipped = 0
    for e, at in plan:
        if at is None:
            skipped += 1
            print(f"skip {e['row_id']}: already corrected")
            continue
        parts = final[e["source_slug"]]
        lead = parts[at][:len(parts[at]) - len(parts[at].lstrip())]
        trail = parts[at][len(parts[at].rstrip()):]
        parts[at] = f"{lead}{e['new_verb']}: {e['new_target']}{trail}"
        changed += 1
        print(f"set  {e['row_id']}: {e['old_verb']}: {e['old_target']} -> {e['new_verb']}: {e['new_target']}")
    for slug, parts in final.items():
        before = len(segments(by_slug[slug][idx["related_concepts"]].value))
        if len(segments("|".join(parts))) != before:
            errors.append(f"{slug}: the corrected cell would have a different number of segments")
    for e, _ in plan:
        n = sum(parse(p) == (e["new_verb"], e["new_target"]) for p in final[e["source_slug"]])
        if n != 1:
            errors.append(f"{e['row_id']}: the corrections together leave '{e['new_verb']}: {e['new_target']}' "
                          f"{n} times in {e['source_slug']}'s cell")
    if errors:
        sys.exit("refused, nothing written:\n  " + "\n  ".join(errors))
    for slug, parts in final.items():
        by_slug[slug][idx["related_concepts"]].value = "|".join(parts)

    print(f"{changed} changed, {skipped} already corrected")
    if changed and not args.check:
        wb.save(args.workbook)
        print(f"saved {args.workbook}")


if __name__ == "__main__":
    main()
