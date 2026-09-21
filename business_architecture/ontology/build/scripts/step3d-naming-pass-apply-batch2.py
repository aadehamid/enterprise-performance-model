#!/usr/bin/env python3
"""Step 3d naming pass - scope-ambiguity Batch 2 (D, 11 renames).

Approved by Hamid 2026-09-21 on independent reviewer recommendation
(approve as proposed, with the KYC scope-boundary condition verified
against the workbook before execution).

The KYC condition is satisfied by the existing approved text: the
CM-1-3-7-4-6 scope note already states "Execution, not framework
ownership: the sanctions compliance and financial-crime frameworks are
owned by Compliance and Legal; this process executes them", that the
process "does not clear a true/false match, grant exceptions, or
determine permissibility of a relationship", and the out_of_scope
column reserves "Sanctions/AML framework ownership, match clearing,
exceptions, and relationship-permissibility determinations" to
Compliance/Legal. No definition change is made in this naming batch.

Treatment (same discipline as Batch 1, PR #91):
  - prefLabel <- new label; the old label is retained as skos:altLabel
    (all 11 legacy labels verified uniquely identifiable, no
    prefLabel collision with another concept).
  - CM-1-3-7-1-6 uses the accepted wording "Maintain Commercial Policy
    Content and Approved Parameters" (the workbook alt column carries
    the variant without "Approved"; the proposal wording wins).
  - identity map: name <- new, prior_name <- old, scoped_historical_alias
    stays null, name_change_note recorded.
  - related_concepts "predicate: Old" references become
    "predicate: New (CM-slug)" (slug annotation, Batch 1 convention).
  - prose columns (definition/scope/in/out/key inputs/outputs/open
    questions/terminology notes) and breadcrumbs: plain old -> new.
  - execution record appended to each renamed row's terminology_notes
    AFTER the global replacements, so the record's own old-label
    mentions survive verbatim.

One-shot migration script. Every replacement asserts its expected count,
so a stale workbook fails loudly instead of silently under-applying.
Run from the repo root:

    python3 build/scripts/step3d-naming-pass-apply-batch2.py

Edits build/output/step3c-definition-authoring-workbook.xlsx,
build/output/step2-identity-map.json, and step3c-naming-pass-queue.md
in place.
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[4]
WB = ROOT / "business_architecture/ontology/build/output/step3c-definition-authoring-workbook.xlsx"
SHEET = "Review & authoring"
IDENTITY_MAP = ROOT / "business_architecture/ontology/build/output/step2-identity-map.json"
QUEUE = ROOT / "business_architecture/ontology/step3c-naming-pass-queue.md"

NEW = {
    "CM-1-3-7-1": "Manage Commercial Master Data Stewardship",
    "CM-1-3-7-1-1": "Maintain Customer and Commercial Account Master Data",
    "CM-1-3-7-1-3": "Maintain Product and Service Master Data",
    "CM-1-3-7-1-5": "Maintain Commercial Workflow Configuration",
    "CM-1-3-7-1-6": "Maintain Commercial Policy Content and Approved Parameters",
    "CM-1-3-7-2": "Manage Commercial Terms, Quoting, and Customer Commercial Services",
    "CM-1-3-7-3-6": "Manage Commercial Returns Authorization and Coordination",
    "CM-1-3-7-4-5": "Perform Credit-Driven Customer Closure and Reinstatement",
    "CM-1-3-7-4-6": "Perform KYC Due Diligence",
    "CM-1-3-7-5-1": "Prioritize Customer Requests and Inquiries",
    "CM-1-3-7-5-2": "Maintain Customer Request and Inquiry Records",
}
OLD = {
    "CM-1-3-7-1": "Manage Master Data (Customer, Price, Product, Tax, Location)",
    "CM-1-3-7-1-1": "Maintain Customer Master Data",
    "CM-1-3-7-1-3": "Maintain Product Master Data",
    "CM-1-3-7-1-5": "Maintain Workflows",
    "CM-1-3-7-1-6": "Maintain Commercial Policies",
    "CM-1-3-7-2": "Manage Pricing & Contracting (Manage Quote & Sale)",
    "CM-1-3-7-3-6": "Manage Returns",
    "CM-1-3-7-4-5": "Perform Customer Closure & Reinstatement",
    "CM-1-3-7-4-6": "KYC Process",
    "CM-1-3-7-5-1": "Establish Priority Assignment of Customer Request/Inquiry",
    "CM-1-3-7-5-2": "Maintain Requests and Inquiries Master Database",
}
# Pre-edit alt_labels cell contents (each row carries exactly the queued
# new label as its alt label; 7-1-6 carries the no-"Approved" variant).
OLD_ALT = {
    "CM-1-3-7-1": "Manage Commercial Master Data Stewardship",
    "CM-1-3-7-1-1": "Maintain Customer and Commercial Account Master Data",
    "CM-1-3-7-1-3": "Maintain Product and Service Master Data",
    "CM-1-3-7-1-5": "Maintain Commercial Workflow Configuration",
    "CM-1-3-7-1-6": "Maintain Commercial Policy Content and Parameters",
    "CM-1-3-7-2": "Manage Commercial Terms, Quoting, and Customer Commercial Services",
    "CM-1-3-7-3-6": "Manage Commercial Returns Authorization and Coordination",
    "CM-1-3-7-4-5": "Perform Credit-Driven Customer Closure and Reinstatement",
    "CM-1-3-7-4-6": "Perform KYC Due Diligence",
    "CM-1-3-7-5-1": "Prioritize Customer Requests and Inquiries",
    "CM-1-3-7-5-2": "Maintain Customer Request and Inquiry Records",
}

PROSE_COLS = (
    "definition", "scope_note", "in_scope", "out_of_scope",
    "key_inputs", "primary_output", "open_questions", "terminology_notes",
)

# (slug, kind, expected_hits) — kinds: "rc" (related_concepts, gets slug
# annotation), "prose" (plain replace across PROSE_COLS), "bc" (breadcrumb).
PLAN = [
    ("CM-1-3-7-1", "rc", 4), ("CM-1-3-7-1", "bc", 7),
    ("CM-1-3-7-1-1", "rc", 7), ("CM-1-3-7-1-1", "prose", 5), ("CM-1-3-7-1-1", "bc", 1),
    ("CM-1-3-7-1-3", "bc", 1),
    ("CM-1-3-7-1-5", "bc", 1),
    ("CM-1-3-7-1-6", "rc", 3), ("CM-1-3-7-1-6", "prose", 1), ("CM-1-3-7-1-6", "bc", 1),
    ("CM-1-3-7-2", "rc", 2), ("CM-1-3-7-2", "bc", 7),
    ("CM-1-3-7-3-6", "rc", 2), ("CM-1-3-7-3-6", "bc", 1),
    ("CM-1-3-7-4-5", "prose", 1), ("CM-1-3-7-4-5", "bc", 1),
    ("CM-1-3-7-4-6", "rc", 4), ("CM-1-3-7-4-6", "prose", 3), ("CM-1-3-7-4-6", "bc", 1),
    ("CM-1-3-7-5-1", "rc", 1), ("CM-1-3-7-5-1", "bc", 1),
    ("CM-1-3-7-5-2", "bc", 1),
]

EXEC_NOTE_TMPL = (
    "Step 3d naming pass \u2014 scope-ambiguity Batch 2 (D, 2026-09-21): "
    "renamed from '{old}' to '{new}'. Prior label retained as altLabel."
)


def main() -> int:
    from openpyxl import load_workbook

    wb = load_workbook(WB)
    ws = wb[SHEET]
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    by_slug = {}
    for r in ws.iter_rows(min_row=2):
        s = r[idx["slug"]].value
        if s:
            by_slug[s] = r

    # 1. the 11 renames (name/alt swap only; terminology notes later)
    for slug, new_name in NEW.items():
        row = by_slug[slug]
        old_name = row[idx["name"]].value
        assert old_name == OLD[slug], f"{slug}: name {old_name!r} != expected {OLD[slug]!r}"
        old_alt = row[idx["alt_labels"]].value or ""
        assert old_alt == OLD_ALT[slug], f"{slug}: alt {old_alt!r} != expected {OLD_ALT[slug]!r}"
        row[idx["name"]].value = new_name
        row[idx["alt_labels"]].value = OLD[slug]
        print(f"renamed {slug}: {OLD[slug]!r} -> {new_name!r}")

    # 2. global replacements with asserted counts
    for slug, kind, expected in PLAN:
        old, new = OLD[slug], NEW[slug]
        esc = re.escape(old)
        total = 0
        if kind == "rc":
            rx = re.compile(r"([a-z][a-z-]*): " + esc + r"(?= \||$)")
            repl = r"\1: " + new + f" ({slug})"
            cols = [idx["related_concepts"]]
        elif kind == "prose":
            rx = re.compile(esc)
            repl = new
            cols = [idx[c] for c in PROSE_COLS]
        else:  # breadcrumb
            rx = re.compile(esc)
            repl = new
            cols = [idx["breadcrumb"]]
        for r in ws.iter_rows(min_row=2):
            for ci in cols:
                v = r[ci].value
                if not v:
                    continue
                new_v, n = rx.subn(repl, v)
                if n:
                    r[ci].value = new_v
                    total += n
        assert total == expected, f"{slug}/{kind}: got {total}, expected {expected}"
        print(f"global {slug}/{kind}: {total} (expected {expected}) OK")

    # 3. execution records (after globals so old-label mentions survive)
    for slug, new_name in NEW.items():
        row = by_slug[slug]
        note = EXEC_NOTE_TMPL.format(old=OLD[slug], new=new_name)
        cur = row[idx["terminology_notes"]].value or ""
        row[idx["terminology_notes"]].value = (cur + " " + note).strip() if cur else note

    # 4. final sweep: no stale references may remain outside the 11
    # renamed rows' own terminology_notes (execution records legitimately
    # mention the old labels), name, and alt_labels (retained aliases).
    stale = []
    check_cols = [idx[c] for c in PROSE_COLS] + [idx["related_concepts"], idx["breadcrumb"]]
    for r in ws.iter_rows(min_row=2):
        s = r[idx["slug"]].value
        for ci in check_cols:
            if s in NEW and ci == idx["terminology_notes"]:
                continue
            v = r[ci].value or ""
            for slug, old in OLD.items():
                if re.search(re.escape(old), v):
                    stale.append((s, headers[ci], old))
    assert not stale, f"stale references remain: {stale[:10]}"
    print("final sweep: no stale references outside the 11 renamed rows' own notes")

    wb.save(WB)
    print(f"saved {WB}")

    # 5. identity map overlays
    entries = json.loads(IDENTITY_MAP.read_text(encoding="utf-8"))
    by = {e["slug"]: e for e in entries}
    for slug, new_name in NEW.items():
        e = by[slug]
        assert e["name"] == OLD[slug], f"identity map {slug}: {e['name']!r} != {OLD[slug]!r}"
        assert e.get("prior_name") in (None, ""), f"identity map {slug} already has prior_name"
        e["name"] = new_name
        e["prior_name"] = OLD[slug]
        e["scoped_historical_alias"] = None
        e["name_change_note"] = (
            "Step 3d naming pass \u2014 scope-ambiguity Batch 2 (D, 2026-09-21): "
            f"renamed '{OLD[slug]}' -> '{new_name}' (Hamid approval 2026-09-21, "
            "independent reviewer recommendation). Prior label retained as "
            "skos:altLabel; no identifier, hierarchy, or definition change."
        )
    IDENTITY_MAP.write_text(json.dumps(entries, indent=1, ensure_ascii=False) + "\n",
                            encoding="utf-8")
    print(f"identity map: 11 overlays updated")

    # 6. naming queue: mark the 11 D rows EXECUTED
    text = QUEUE.read_text(encoding="utf-8")
    lines = text.split("\n")
    marked = 0
    for i, line in enumerate(lines):
        for slug, new_name in NEW.items():
            if line.startswith(f"| `{slug}` |") and "**EXECUTED" not in line:
                assert line.rstrip().endswith("|"), f"queue row has no trailing pipe: {line[:60]}"
                lines[i] = line.rstrip()[:-1] + (
                    f"  **EXECUTED 2026-09-21** (Step 3d naming pass, scope-ambiguity "
                    f"Batch 2 D: renamed to '{new_name}'; prior label retained as "
                    f"altLabel; see naming PR). |"
                )
                marked += 1
                break
    assert marked == 11, f"queue rows marked: {marked}, expected 11"
    text = "\n".join(lines)
    old_hdr = "## Scope ambiguity (33 \u2014 10 executed 2026-09-21, 23 remaining)"
    new_hdr = "## Scope ambiguity (33 \u2014 21 executed 2026-09-21, 12 remaining)"
    assert old_hdr in text, "queue scope-ambiguity header not in expected form"
    text = text.replace(old_hdr, new_hdr)
    old_status = ("The remaining 78 queue entries are still queued for the consolidated naming pass.")
    new_status = ("10 scope-ambiguity entries executed 2026-09-21 (Step 3d naming pass Batch 1 A+B+C, "
                  "Hamid approval 2026-09-21) and 11 scope-ambiguity entries executed 2026-09-21 "
                  "(Step 3d naming pass Batch 2 D, Hamid approval 2026-09-21). The remaining 57 "
                  "queue entries are still queued for the consolidated naming pass.")
    assert old_status in text, "queue execution-status line not in expected form"
    text = text.replace(old_status, new_status)
    QUEUE.write_text(text, encoding="utf-8")
    print("naming queue: 11 D rows marked EXECUTED")

    return 0


if __name__ == "__main__":
    sys.exit(main())
