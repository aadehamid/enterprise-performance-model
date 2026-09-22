#!/usr/bin/env python3
"""Step 3d naming pass - generic-operational Batch 5 (11 renames).

Approved by Hamid 2026-09-21 (explicit "Go ahead" in chat on the reviewer's
verdict: approve with one label refinement and one alias-control exception).

The batch converts labels that depended on parent context into standalone,
searchable, semantically safe labels. Three subclusters:
  A. Terminal Commercial Operations (CM-1-3-10-1..4): terminal-specific
     record/authorization, lifting forecasts & nominations, lifting
     allocation, and terminal sales-deal capture.
  B. Commercial Agreement Compliance (CM-1-3-6-8): narrows the over-broad
     "Compliance Management" to its approved commercial-agreement scope.
  C. Service & Support cluster (CM-1-3-7-5-5, CM-1-3-9-1-2, CM-1-3-9-2-1,
     CM-1-3-9-2-2, CM-1-3-9-2-3, CM-1-3-9-3-2): service operations, service
     delivery data, customer interactions, service events, customer service
     reviews, service workforce performance.

Refinements vs the queue (per the approved verdict):
  1. CM-1-3-9-2-2 executes the PLURAL "Manage Customer Interactions in
     Service Delivery" (queued alt was the singular "Manage Customer
     Interaction in Service Delivery").
  2. Alias-control exception for CM-1-3-9-2-2: the bare historical label
     "Manage Customer" is NOT retained as skos:altLabel. The exact-collision
     check found no other row with it as pref/alt label, but the bare label
     is a prefix of live labels (Manage Customer Portal, Manage Customer
     Invoicing and Billing, Manage Customer Requests and Inquiries) and
     would reintroduce the retrieval ambiguity this batch removes. It is
     preserved as prior_name in the identity/migration map in all cases.

Guard patterns (three new labels contain their old label as a substring):
  - "Commercial Agreement Compliance Management":  (?<!Agreement )Compliance Management
  - "Manage Customer Interactions in Service Delivery": Manage Customer (standalone
    only, followed by " |" or end — never the longer "Manage Customer *"
    live labels, never the new plural label)
  - "Fulfill Service Events": Fulfill Service Event not followed by "s"

Treatment (same discipline as Batches 1-4):
  - prefLabel <- new label; the old label is retained as skos:altLabel
    (except "Manage Customer" per the exception above).
  - identity map: name <- new, prior_name <- old,
    scoped_historical_alias stays null, name_change_note recorded.
  - related_concepts "predicate: Old" references become
    "predicate: New (CM-slug)" (slug annotation, Batch 1 convention).
  - prose columns, breadcrumbs, and the parent-display cell updated.
  - execution record appended to each renamed row's terminology_notes
    AFTER the global replacements, so the record's own old-label
    mentions survive verbatim.

One-shot migration script. Every replacement asserts its expected count,
so a stale workbook fails loudly instead of silently under-applying.
Run from the repo root:

    python3 business_architecture/ontology/build/scripts/step3d-naming-pass-apply-batch5.py

Edits build/output/step3c-definition-authoring-workbook.xlsx,
build/output/step2-identity-map.json, and step3c-naming-pass-queue.md
in place.
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[4]
WB = ROOT / "business_architecture/ontology/build/output/step3c-definition-authoring-workbook.xlsx"
SHEET = "Review & authoring"
IDENTITY_MAP = ROOT / "business_architecture/ontology/build/output/step2-identity-map.json"
QUEUE = ROOT / "business_architecture/ontology/step3c-naming-pass-queue.md"

NEW = {
    "CM-1-3-10-1": "Set Up and Maintain Terminal Customer Authorization",
    "CM-1-3-10-2": "Process Customer Lifting Forecasts and Nominations",
    "CM-1-3-10-3": "Manage Terminal Lifting Allocation",
    "CM-1-3-10-4": "Capture Terminal Sales Deal",
    "CM-1-3-6-8": "Commercial Agreement Compliance Management",
    "CM-1-3-7-5-5": "Conduct Customer Service Reviews",
    "CM-1-3-9-1-2": "Manage Service Operations",
    "CM-1-3-9-2-1": "Manage Service Delivery Data",
    # plural refinement per the approved verdict (queued alt was singular)
    "CM-1-3-9-2-2": "Manage Customer Interactions in Service Delivery",
    "CM-1-3-9-2-3": "Fulfill Service Events",
    "CM-1-3-9-3-2": "Measure Service Workforce Performance",
}
OLD = {
    "CM-1-3-10-1": "Setup and Maintain Customer In Terminal",
    "CM-1-3-10-2": "Process Forecast and Nominations",
    "CM-1-3-10-3": "Manage Allocation",
    "CM-1-3-10-4": "Capture Deal",
    "CM-1-3-6-8": "Compliance Management",
    "CM-1-3-7-5-5": "Conduct Quarterly Review Meeting",
    "CM-1-3-9-1-2": "Manage Operations",
    "CM-1-3-9-2-1": "Manage Data",
    "CM-1-3-9-2-2": "Manage Customer",
    "CM-1-3-9-2-3": "Fulfill Service Event",
    "CM-1-3-9-3-2": "Measure Service Employees",
}
# workbook queued alt_labels (asserted before the swap); singular for 9-2-2
QUEUED_ALT = {
    "CM-1-3-10-1": "Set Up and Maintain Terminal Customer Authorization",
    "CM-1-3-10-2": "Process Customer Lifting Forecasts and Nominations",
    "CM-1-3-10-3": "Manage Terminal Lifting Allocation",
    "CM-1-3-10-4": "Capture Terminal Sales Deal",
    "CM-1-3-6-8": "Commercial Agreement Compliance Management",
    "CM-1-3-7-5-5": "Conduct Customer Service Reviews",
    "CM-1-3-9-1-2": "Manage Service Operations",
    "CM-1-3-9-2-1": "Manage Service Delivery Data",
    "CM-1-3-9-2-2": "Manage Customer Interaction in Service Delivery",
    "CM-1-3-9-2-3": "Fulfill Service Events",
    "CM-1-3-9-3-2": "Measure Service Workforce Performance",
}

# slug of the alias-control exception (historical label NOT kept as altLabel)
ALIAS_EXCEPTION = "CM-1-3-9-2-2"

PROSE_COLS = (
    "definition", "scope_note", "in_scope", "out_of_scope",
    "key_inputs", "primary_output", "open_questions", "terminology_notes",
)

# Guarded old-label patterns: match the standalone old label but never the
# new label (substring case) and never longer live labels sharing the prefix.
def old_pat(slug: str) -> str:
    old = re.escape(OLD[slug])
    if slug == "CM-1-3-6-8":
        return r"(?<!Agreement )" + old
    if slug == "CM-1-3-9-2-3":
        return old + r"(?!s)"
    return old

# (slug, kind, expected_hits)
# kinds: "rc" (related_concepts, gets slug annotation), "prose" (plain
# replace across PROSE_COLS), "bc" (breadcrumb), "parent" (parent-display).
PLAN = [
    ("CM-1-3-10-1", "rc", 3), ("CM-1-3-10-1", "bc", 1),
    ("CM-1-3-10-2", "rc", 3), ("CM-1-3-10-2", "bc", 1),
    ("CM-1-3-10-3", "rc", 2), ("CM-1-3-10-3", "prose", 3), ("CM-1-3-10-3", "bc", 1),
    ("CM-1-3-10-4", "rc", 2), ("CM-1-3-10-4", "prose", 2), ("CM-1-3-10-4", "bc", 1),
    ("CM-1-3-6-8", "rc", 1), ("CM-1-3-6-8", "bc", 2), ("CM-1-3-6-8", "parent", 1),
    ("CM-1-3-7-5-5", "rc", 1), ("CM-1-3-7-5-5", "bc", 1),
    ("CM-1-3-9-1-2", "rc", 5), ("CM-1-3-9-1-2", "bc", 1),
    ("CM-1-3-9-2-1", "bc", 1),
    ("CM-1-3-9-2-2", "rc", 1), ("CM-1-3-9-2-2", "bc", 1),
    ("CM-1-3-9-2-3", "rc", 3), ("CM-1-3-9-2-3", "bc", 1),
    ("CM-1-3-9-3-2", "bc", 1),
]

EXEC_NOTE_TMPL = (
    "Step 3d naming pass \u2014 generic-operational Batch 5 (2026-09-21): "
    "renamed from '{old}' to '{new}'. {alias_note}"
)
EXEC_NOTE_ALIAS_STD = "Prior label retained as altLabel."
EXEC_NOTE_ALIAS_EXC = (
    "Prior label 'Manage Customer' NOT retained as skos:altLabel per the "
    "alias-control exception (bare label is a prefix of live labels and "
    "would reintroduce retrieval ambiguity); preserved as prior_name in "
    "the identity/migration map."
)


def main() -> int:
    from openpyxl import load_workbook

    wb = load_workbook(WB)
    ws = wb[SHEET]
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    by_slug = {}
    for r in ws.iter_rows(min_row=2):
        s = r[idx["slug"]].value
        if s:
            by_slug[s] = r

    # 1. the 11 renames (name/alt swap; terminology notes later)
    for slug, new_name in NEW.items():
        row = by_slug[slug]
        old_name = row[idx["name"]].value
        assert old_name == OLD[slug], f"{slug}: name {old_name!r} != expected {OLD[slug]!r}"
        old_alt = row[idx["alt_labels"]].value or ""
        assert old_alt == QUEUED_ALT[slug], f"{slug}: alt {old_alt!r} != expected {QUEUED_ALT[slug]!r}"
        row[idx["name"]].value = new_name
        if slug == ALIAS_EXCEPTION:
            row[idx["alt_labels"]].value = ""
            print(f"renamed {slug}: {OLD[slug]!r} -> {new_name!r} (alias exception: no altLabel)")
        else:
            row[idx["alt_labels"]].value = OLD[slug]
            print(f"renamed {slug}: {OLD[slug]!r} -> {new_name!r}")

    # 2. scoped replacements with asserted counts
    for slug, kind, expected in PLAN:
        old, new = OLD[slug], NEW[slug]
        pat = old_pat(slug)
        total = 0
        if kind == "rc":
            if slug == "CM-1-3-9-2-2":
                # standalone label only: never the longer "Manage Customer *"
                # live labels, never the new plural label
                rx = re.compile(r"([a-z][a-z-]*): Manage Customer(?= \||$)")
            else:
                rx = re.compile(r"([a-z][a-z-]*): " + pat + r"(?= \||$)")
            repl = r"\1: " + new + f" ({slug})"
            cols = [idx["related_concepts"]]
        elif kind == "prose":
            rx = re.compile(pat)
            repl = new
            cols = [idx[c] for c in PROSE_COLS]
        elif kind == "parent":
            rx = re.compile(r"^" + pat + r"$")
            repl = new
            cols = [idx["parent"]]
        else:  # breadcrumb
            if slug == "CM-1-3-9-2-2":
                rx = re.compile(r"Manage Customer$")
            else:
                rx = re.compile(pat)
            repl = new
            cols = [idx["breadcrumb"]]
        for r in ws.iter_rows(min_row=2):
            for ci in cols:
                v = r[ci].value
                if not v:
                    continue
                new_v, n = rx.subn(repl, v)
                if n:
                    r[ci].value = new_v
                    total += n
        assert total == expected, f"{slug}/{kind}: got {total}, expected {expected}"
        print(f"{slug}/{kind}: {total} (expected {expected}) OK")

    # 3. execution records (after globals so old-label mentions survive)
    for slug, new_name in NEW.items():
        row = by_slug[slug]
        alias_note = EXEC_NOTE_ALIAS_EXC if slug == ALIAS_EXCEPTION else EXEC_NOTE_ALIAS_STD
        note = EXEC_NOTE_TMPL.format(old=OLD[slug], new=new_name, alias_note=alias_note)
        cur = row[idx["terminology_notes"]].value or ""
        row[idx["terminology_notes"]].value = (cur + " " + note).strip() if cur else note

    # 4. final sweep: no stale references may remain outside —
    #    (a) the 11 renamed rows' own terminology_notes (execution records
    #        legitimately mention the old labels),
    #    (b) name and alt_labels (retained aliases; row 9-2-2 keeps none).
    #    Guarded patterns so longer live labels ("Manage Customer Portal",
    #    "Manage Customer Invoicing and Billing", "Manage Customer Requests
    #    and Inquiries") and the new labels themselves never flag.
    def sweep_pat(slug: str) -> str:
        if slug == "CM-1-3-9-2-2":
            return r"Manage Customer(?= \||$|>)"
        return old_pat(slug)

    stale = []
    check_cols = [idx[c] for c in PROSE_COLS] + [idx["related_concepts"], idx["breadcrumb"], idx["parent"]]
    for r in ws.iter_rows(min_row=2):
        s = r[idx["slug"]].value
        for ci in check_cols:
            if s in NEW and ci == idx["terminology_notes"]:
                continue
            if s in NEW and ci in (idx["name"], idx["alt_labels"]):
                continue
            v = r[ci].value or ""
            for slug in OLD:
                if re.search(sweep_pat(slug), v):
                    stale.append((s, headers[ci], OLD[slug]))
    assert not stale, f"stale references remain: {stale[:10]}"
    print("final sweep: no stale references outside renamed rows' notes/aliases")

    wb.save(WB)
    print(f"saved {WB}")

    # 5. identity map overlays
    entries = json.loads(IDENTITY_MAP.read_text(encoding="utf-8"))
    by = {e["slug"]: e for e in entries}
    for slug, new_name in NEW.items():
        e = by[slug]
        assert e["name"] == OLD[slug], f"identity map {slug}: {e['name']!r} != {OLD[slug]!r}"
        assert e.get("prior_name") in (None, ""), f"identity map {slug} already has prior_name"
        e["name"] = new_name
        e["prior_name"] = OLD[slug]
        e["scoped_historical_alias"] = None
        if slug == ALIAS_EXCEPTION:
            e["name_change_note"] = (
                "Step 3d naming pass \u2014 generic-operational Batch 5 (2026-09-21): "
                f"renamed '{OLD[slug]}' -> '{new_name}' (Hamid approval 2026-09-21; "
                "independent reviewer recommendation: approve with plural "
                "'Interactions' refinement for the enduring multi-interaction "
                "process). Alias-control exception: the bare historical label "
                "'Manage Customer' is NOT retained as skos:altLabel \u2014 "
                "uniqueness/collision validation found no exact pref/alt "
                "collision on other rows, but the bare label is a prefix of "
                "live labels (Manage Customer Portal, Manage Customer "
                "Invoicing and Billing, Manage Customer Requests and "
                "Inquiries) and would reintroduce retrieval ambiguity; it is "
                "preserved as prior_name in the identity/migration map only. "
                "No identifier, hierarchy, or definition change."
            )
        else:
            e["name_change_note"] = (
                "Step 3d naming pass \u2014 generic-operational Batch 5 (2026-09-21): "
                f"renamed '{OLD[slug]}' -> '{new_name}' (Hamid approval 2026-09-21; "
                "independent reviewer recommendation: approve as proposed). Prior "
                "label retained as skos:altLabel; no identifier, hierarchy, or "
                "definition change."
            )
    IDENTITY_MAP.write_text(json.dumps(entries, indent=1, ensure_ascii=False) + "\n",
                            encoding="utf-8")
    print("identity map: 11 overlays updated")

    # 6. naming queue: mark the 11 generic-operational rows EXECUTED
    queue_slugs = set(NEW)
    text = QUEUE.read_text(encoding="utf-8")
    lines = text.split("\n")
    marked = 0
    for i, line in enumerate(lines):
        for slug in queue_slugs:
            if line.startswith(f"| `{slug}` |") and "**EXECUTED" not in line:
                assert line.rstrip().endswith("|"), f"queue row has no trailing pipe: {line[:60]}"
                if slug == ALIAS_EXCEPTION:
                    exec_rec = (
                        f"  **EXECUTED 2026-09-21** (Step 3d naming pass, "
                        f"generic-operational Batch 5: renamed to '{NEW[slug]}' "
                        f"(plural refinement of the queued singular label); "
                        f"historical label '{OLD[slug]}' NOT retained as "
                        f"altLabel per the alias-control exception, kept as "
                        f"prior_name in the identity map; see naming PR). |"
                    )
                else:
                    exec_rec = (
                        f"  **EXECUTED 2026-09-21** (Step 3d naming pass, "
                        f"generic-operational Batch 5: renamed to '{NEW[slug]}'; "
                        f"prior label retained as altLabel; see naming PR). |"
                    )
                lines[i] = line.rstrip()[:-1] + exec_rec
                marked += 1
                break
    assert marked == 11, f"queue rows marked: {marked}, expected 11"
    text = "\n".join(lines)
    old_hdr = "## Generic operational label (11)"
    new_hdr = "## Generic operational label (11 \u2014 11 executed 2026-09-21, 0 remaining)"
    assert old_hdr in text, "queue generic-operational header not in expected form"
    text = text.replace(old_hdr, new_hdr)
    old_status = ("and 8 directionality-missing entries executed 2026-09-21 (Step 3d naming pass "
                  "Batch 4, Hamid approval 2026-09-21). The remaining 37 queue entries "
                  "are still queued for the consolidated naming pass.")
    new_status = ("and 8 directionality-missing entries executed 2026-09-21 (Step 3d naming pass "
                  "Batch 4, Hamid approval 2026-09-21) and 11 generic-operational "
                  "entries executed 2026-09-21 (Step 3d naming pass Batch 5, Hamid approval "
                  "2026-09-21). The remaining 26 queue entries are still queued for the "
                  "consolidated naming pass.")
    assert old_status in text, "queue execution-status line not in expected form"
    text = text.replace(old_status, new_status)
    QUEUE.write_text(text, encoding="utf-8")
    print("naming queue: 11 generic-operational rows marked EXECUTED")

    return 0


if __name__ == "__main__":
    sys.exit(main())
