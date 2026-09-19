#!/usr/bin/env python3
"""Step 3c workbook review gate (mechanical checks + review-findings report).

Usage:
  python3 step3c-workbook-validate.py --workbook <xlsx> [--ptc-register <md>]
      [--identity-map <json>] [--out findings.md]

Reads the 'Review & authoring' sheet, validates every row, cross-checks the
Parked Tree Changes register, and writes a findings report.

Statuses: pending | approved | blocked | retired
  - approved: definition + scope_note required; merges into the taxonomy.
  - blocked:  row cannot be defined until its PTC entry closes; must cite
               the PTC ID in terminology_notes; must NOT carry a definition.
  - retired:  row will not be defined (node leaves/is removed in tree pass);
               must cite the PTC ID; must NOT carry a definition.
  - pending:  queued for authoring.

PTC register enforcement (the register is authoritative on tree changes):
  - blocked/retired row naming no PTC            -> BLOCKING
  - row citing a PTC ID absent from the register -> BLOCKING
  - blocked row parked against a Closed entry    -> BLOCKING
    (retired may keep the citation as provenance)
  - Open entry cited by no row                   -> BLOCKING
  - zero pending rows while any entry is Open    -> BLOCKING (step-3c-not-complete)

Exit code 0 = no blocking findings; 1 = blocking findings present.
"""
import argparse, json, os, re, sys
from openpyxl import load_workbook

BLOCKING, QUESTION, NOTE = "BLOCKING", "QUESTION", "NOTE"
VALID_STATUSES = ("pending", "approved", "blocked", "retired")


def parse_ptc_register(path):
    """Return {ptc_id: 'open'|'closed'} parsed from the register markdown."""
    entries = {}
    if not path:
        return entries
    text = open(path).read()
    for m in re.finditer(r"^##\s+(PTC-\d+)\b(.*?)(?=^##\s+PTC-\d+|\Z)",
                         text, re.M | re.S):
        pid, body = m.group(1), m.group(2)
        sm = re.search(r"^\*\*Status:\*\*\s*(.+)$", body, re.M)
        status_line = sm.group(1) if sm else ""
        entries[pid] = "closed" if "closed" in status_line.lower() else "open"
    return entries


# The 15 rows approved before the Phase-1 intake columns existed (#29).
# Flagged as NOTE when Phase 1 is empty — not silently waived, not blocked,
# not backfilled here.
PRE_INTAKE_APPROVED = {
    "CM-1", "CM-1-1-2", "CM-1-1-4", "CM-1-1-5", "CM-1-1-6", "CM-1-3-3-5-6",
    "L1-refining", "L1-midstream", "L1-supply-chain-mgmt", "L1-finance",
    "L1-shared-services", "L1-process-excellence-it", "L1-human-resources",
    "L1-legal-corp-comm", "L1-ehs-gov-reporting",
}


def controlled_values(wb, field):
    """Read allowed values for a field from the 'Controlled vocabularies' sheet."""
    if "Controlled vocabularies" not in wb.sheetnames:
        return set()
    for row in wb["Controlled vocabularies"].iter_rows(min_row=3,
                                                       values_only=True):
        if row[0] == field and row[1]:
            return {v.strip() for v in str(row[1]).split("|")}
    return set()


def registered_sources(wb):
    """Return the set of source_id values from the 'Reference register' sheet."""
    ids = set()
    if "Reference register" not in wb.sheetnames:
        return ids
    rows = list(wb["Reference register"].iter_rows(values_only=True))
    start = next((i for i, r in enumerate(rows)
                  if any("source_id" in str(c or "") for c in r)), None)
    if start is None:
        return ids
    for r in rows[start + 1:]:
        sid = str(r[0] or "").strip()
        if sid and sid != "source_id":
            ids.add(sid)
    return ids


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--workbook", required=True)
    ap.add_argument("--ptc-register",
                    default=os.path.join(os.path.dirname(
                        os.path.abspath(__file__)),
                        "..", "..", "step3c-parked-tree-changes.md"),
                    help="PTC register markdown ('' to skip register checks)")
    ap.add_argument("--identity-map",
                    default=os.path.join(os.path.dirname(
                        os.path.abspath(__file__)),
                        "..", "output", "step2-identity-map.json"),
                    help="Step 2 identity map JSON (preferred-label source)")
    ap.add_argument("--out", default="step3c-workbook-findings.md")
    a = ap.parse_args()

    ptc = parse_ptc_register(a.ptc_register if a.ptc_register else None)

    wb = load_workbook(a.workbook, data_only=True)
    ws = wb["Review & authoring"]
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}

    findings, seen_slugs = [], set()
    cited_ptc = set()
    stats = {"rows": 0, "approved": 0, "pending": 0, "blocked": 0, "retired": 0}

    def add(sev, slug, check, detail):
        findings.append({"severity": sev, "slug": slug, "check": check,
                         "detail": detail})

    try:
        idmap = {r["slug"]: r for r in json.load(open(a.identity_map))}
    except FileNotFoundError:
        idmap = {}
    pref_labels = {r["name"].lower() for r in idmap.values()}
    CTC_ALLOWED = controlled_values(wb, "concept_type_check") or \
        {"process", "capability", "mixed/needs-review", "not-process"}
    HORIZONS = controlled_values(wb, "process_horizon")
    REGISTERED = registered_sources(wb)

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        g = lambda h: (str(row[idx[h]]) if h in idx and idx[h] < len(row)
                       and row[idx[h]] is not None else "").strip()
        slug = g("slug")
        stats["rows"] += 1
        status = g("status").lower()
        defi = g("definition")
        scope = g("scope_note")
        out_sc = g("out_of_scope")
        alt = g("alt_labels")
        tnotes = g("terminology_notes")

        if slug in seen_slugs:
            add(BLOCKING, slug, "duplicate-slug", "Slug appears more than once.")
        seen_slugs.add(slug)
        if status not in VALID_STATUSES:
            add(BLOCKING, slug, "bad-status",
                f"status={status!r}; must be one of {list(VALID_STATUSES)}.")
            continue
        stats[status] += 1

        ptc_refs = set(re.findall(r"PTC-\d+", tnotes))
        cited_ptc |= ptc_refs

        if status in ("blocked", "retired"):
            if not ptc_refs:
                add(BLOCKING, slug, "ptc-not-cited",
                    f"status={status!r} but terminology_notes names no PTC entry.")
            for pid in ptc_refs:
                if pid not in ptc:
                    add(BLOCKING, slug, "ptc-unknown",
                        f"Cites {pid}, which has no entry in the PTC register.")
                elif ptc[pid] == "closed" and status == "blocked":
                    add(BLOCKING, slug, "ptc-closed",
                        f"Still blocked against {pid}, which is recorded Closed — "
                        f"lift to pending (closure and re-statusing are one change).")
            if defi:
                add(QUESTION, slug, "defined-but-parked",
                    f"status={status!r} yet a definition is present — "
                    f"{'retired rows are not defined' if status == 'retired' else 'blocked rows cannot be defined until unblocked'}.")
            continue

        if status == "pending":
            if defi or scope:
                add(QUESTION, slug, "pending-with-content",
                    "Row is pending but has definition/scope text — intended?")
            continue

        # --- approved-row checks ---
        if not defi:
            add(BLOCKING, slug, "empty-definition",
                "Approved row has no definition.")
            continue
        words = defi.split()
        if len(words) < 8:
            add(QUESTION, slug, "short-definition",
                f"Only {len(words)} words — likely too thin for enterprise grade.")
        # true label-restatement: "X is the X ..." / "X refers to ..." openings
        name_words = re.findall(r"[A-Za-z]+", g("name").lower())
        lead = " ".join(name_words[:4])
        if lead and re.match(
                rf"^(the\s+)?{re.escape(lead)}\s+(is|are|refers?\s+to|means?)\b",
                defi.lower()):
            add(QUESTION, slug, "circular-opening",
                f"Definition restates the label ({g('name')!r}) instead of "
                f"defining the activity — restate by purpose and outcome.")
        if not scope:
            add(BLOCKING, slug, "missing-scope-note",
                "Locked instruction: every approved row requires a non-empty "
                "scope_note with at least one meaningful boundary.")
        if out_sc and not re.search(r"\b(owned by|belongs to|sibling|see |under )\b",
                                    out_sc, re.I):
            add(NOTE, slug, "out-of-scope-owner",
                "Out-of-scope text doesn't name the owning sibling — fine when "
                "ownership isn't established; otherwise consider adding it.")
        for label in [x.strip() for x in alt.split("|") if x.strip()]:
            if label.lower() in pref_labels:
                add(QUESTION, slug, "altlabel-collision",
                    f"altLabel {label!r} collides with another concept's "
                    f"preferred label.")

        # Phase-1 authoring columns (required on new approvals; the 15
        # baseline rows predate them and are grandfathered)
        ctc, pp, refs = g("concept_type_check"), g("primary_purpose"), \
            g("reference_sources")
        if g("open_questions"):
            add(BLOCKING, slug, "open-questions-on-approved",
                "Approved row still has open_questions content — approval "
                "requires no material open question.")
        if slug in PRE_INTAKE_APPROVED:
            for field, check in (("concept_type_check", "phase1-not-backfilled"),
                                 ("primary_purpose", "phase1-not-backfilled"),
                                 ("reference_sources", "phase1-not-backfilled")):
                if not g(field):
                    add(NOTE, slug, check,
                        f"Pre-intake approval (PR #29): {field} empty — "
                        f"flagged, not waived; backfill is a separate decision.")
        if slug not in PRE_INTAKE_APPROVED:
            if not ctc:
                add(BLOCKING, slug, "missing-concept-type",
                    "concept_type_check is required on new approvals.")
            elif ctc not in CTC_ALLOWED:
                add(BLOCKING, slug, "bad-concept-type",
                    f"concept_type_check={ctc!r} is not a controlled value "
                    f"{sorted(CTC_ALLOWED)}.")
            elif ctc not in ("process", "capability"):
                add(BLOCKING, slug, "concept-type-not-approvable",
                    f"concept_type_check={ctc!r}: only 'process' or "
                    f"'capability' can be approved (controlled vocabulary "
                    f"governance).")
            if not pp:
                add(BLOCKING, slug, "missing-primary-purpose",
                    "primary_purpose is required on new approvals.")
            if not refs:
                add(BLOCKING, slug, "missing-reference-sources",
                    "reference_sources is required on new approvals — "
                    "register the source first, then cite its source_id.")
            else:
                for sid in [s.strip() for s in refs.split("|") if s.strip()]:
                    if REGISTERED and sid not in REGISTERED:
                        add(BLOCKING, slug, "unknown-source",
                            f"reference_sources cites {sid!r}, which is not "
                            f"in the Reference register.")
        ph = g("process_horizon")
        if ph and HORIZONS and ph not in HORIZONS:
            add(BLOCKING, slug, "bad-horizon",
                f"process_horizon={ph!r} is not a controlled value "
                f"{sorted(HORIZONS)} — new values need a documented decision.")

    # cross-row register checks
    for pid, state in ptc.items():
        if state == "open" and pid not in cited_ptc:
            add(BLOCKING, pid, "ptc-uncited",
                f"PTC entry is Open but no workbook row cites it — register and "
                f"workbook have drifted.")
    if stats["pending"] == 0 and any(s == "open" for s in ptc.values()):
        add(BLOCKING, "—", "step-3c-not-complete",
            "Definition queue is empty while a PTC entry is still Open — "
            "Step 3c cannot be declared finished with a tree change outstanding.")

    with open(a.out, "w") as f:
        f.write("# Step 3c workbook review findings\n\n")
        f.write(f"Rows: {stats['rows']} | approved: {stats['approved']} | "
                f"pending: {stats['pending']} | blocked: {stats['blocked']} | "
                f"retired: {stats['retired']} | "
                f"open PTC entries: {sum(1 for s in ptc.values() if s == 'open')}\n\n")
        for sev in (BLOCKING, QUESTION, NOTE):
            items = [x for x in findings if x["severity"] == sev]
            f.write(f"## {sev} ({len(items)})\n\n")
            for x in items:
                f.write(f"- **{x['slug']}** [{x['check']}] {x['detail']}\n")
            f.write("\n")
    blocking = sum(1 for x in findings if x["severity"] == BLOCKING)
    print(f"rows={stats['rows']} approved={stats['approved']} "
          f"pending={stats['pending']} blocked={stats['blocked']} "
          f"retired={stats['retired']} "
          f"open_ptc={sum(1 for s in ptc.values() if s == 'open')} "
          f"blocking={blocking} "
          f"questions={sum(1 for x in findings if x['severity'] == QUESTION)}")
    print("report:", a.out)
    return 1 if blocking else 0


if __name__ == "__main__":
    sys.exit(main())
