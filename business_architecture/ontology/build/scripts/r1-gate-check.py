#!/usr/bin/env python3
"""R1 gate checker (PR #108 — Refining structural reclassification).

Independent re-verification of the four R1 merge gates against the
regenerated outputs. Exit 0 = all pass; any FAIL exits 1.

Usage (paths relative to the repository root):
    python3 business_architecture/ontology/build/scripts/r1-gate-check.py \
        [stage_dir] [baseline_ttl] [baseline_identity_map] [workbook]

  stage_dir             Repo tree to check (default: the repo root this
                        script lives in).
  baseline_ttl        The main-branch step3-taxonomy.ttl to diff against
                        (required). Fetch from main via the GitHub API.
  baseline_identity_map The main-branch step2-identity-map.json (required).
  workbook            Step 3c workbook (default: the staged
                        build/output/step3c-definition-authoring-workbook.xlsx).

Requires: rdflib, openpyxl.
"""
import csv
import importlib.util
import json
import sys
from pathlib import Path

from rdflib import Graph, Literal, URIRef
from rdflib.namespace import DCTERMS, OWL, RDF, SKOS  # noqa: F401

REPO_ROOT = Path(__file__).resolve().parents[4]
STAGE = Path(sys.argv[1]) if len(sys.argv) > 1 else REPO_ROOT
if len(sys.argv) < 4:
    sys.exit(__doc__ + "\nERROR: baseline_ttl and baseline_identity_map are required.")
BASE_TTL = Path(sys.argv[2])
BASE_IM = Path(sys.argv[3])
WORKBOOK = (Path(sys.argv[4]) if len(sys.argv) > 4
            else STAGE / "business_architecture/ontology/build/output"
                         "/step3c-definition-authoring-workbook.xlsx")

OUT = STAGE / "business_architecture/ontology/build/output"
R1 = STAGE / "business_architecture/ontology/build/r1"
GEN = STAGE / "business_architecture/ontology/build/step3-skos-taxonomy.py"
PROC = "https://w3id.org/lsc/ontology/process/"
INTAKE = "https://w3id.org/lsc/ontology/intake/"

# Single source of truth for reviewed pending-row interfaces.
spec = importlib.util.spec_from_file_location("step3_skos_taxonomy", GEN)
_gen = importlib.util.module_from_spec(spec)
spec.loader.exec_module(_gen)
REVIEWED_PENDING_INTERFACES = _gen.REVIEWED_PENDING_INTERFACES

results = []


def check(name, cond, detail=""):
    results.append((name, bool(cond), detail))


def load_ttl(p):
    g = Graph()
    g.parse(p, format="turtle")
    return g


g_new = load_ttl(OUT / "step3-taxonomy.ttl")
g_base = load_ttl(BASE_TTL)
im_new = {r["slug"]: r for r in json.load(open(OUT / "step2-identity-map.json"))}
im_base = {r["slug"]: r for r in json.load(open(BASE_IM))}


def concepts(g):
    return set(g.subjects(RDF.type, SKOS.Concept))


C_new, C_base = concepts(g_new), concepts(g_base)


def slug(u):
    return str(u).split("/process/")[-1]


def broader(g, c):
    b = list(g.objects(c, SKOS.broader))
    return b[0] if b else None


def pref(g, c):
    return str(list(g.objects(c, SKOS.prefLabel))[0])


# ---- Gate 1: counts ----
check("concepts 683", len(C_new) == 683, f"{len(C_new)}")
# 14,495 = baseline 14,481 + 14 (new R&T L2); +1 = reviewed interface
# triple on CM-1-1-2-9-1 (Candidate package §1f).
check("triples 14496", len(g_new) == 14496, f"{len(g_new)}")
n_b = len(list(g_new.subject_objects(SKOS.broader)))
check("broader links 681", n_b == 681, f"{n_b}")
roots = [c for c in C_new if broader(g_new, c) is None]
check("exactly 2 roots", len(roots) == 2,
      ", ".join(sorted(slug(r) for r in roots)))
one_broader = all(len(list(g_new.objects(c, SKOS.broader))) == 1
                  for c in C_new if broader(g_new, c) is not None)
check("every non-root has exactly one broader", one_broader)

# ---- dangling + orphans ----
dangling = [(slug(c), slug(p)) for c in C_new
            for p in g_new.objects(c, SKOS.broader) if p not in C_new]
check("0 dangling broader", not dangling, str(dangling[:5]))
check("0 orphan active", len(roots) == 2)

# ---- migration set parents/levels (independent spec) ----
expected = {"CM-1-1-4": ("L1-refining", 2), "CM-1-1-7": ("L1-refining", 2)}
for s in ["CM-1-1-4-1", "CM-1-1-4-2", "CM-1-1-4-3", "CM-1-1-4-4",
          "CM-1-1-4-5", "CM-1-1-4-7"]:
    expected[s] = ("CM-1-1-4", 3)
for s in ["CM-1-1-7-1", "CM-1-1-7-2", "CM-1-1-7-3"]:
    expected[s] = ("CM-1-1-7", 3)
for child in ["CM-1-1-4-1-1", "CM-1-1-4-1-2", "CM-1-1-4-1-3", "CM-1-1-4-1-4",
              "CM-1-1-4-2-1", "CM-1-1-4-2-2",
              "CM-1-1-4-4-1", "CM-1-1-4-4-2", "CM-1-1-4-4-3",
              "CM-1-1-4-7-1", "CM-1-1-4-7-2", "CM-1-1-4-7-3", "CM-1-1-4-7-4",
              "CM-1-1-4-7-5", "CM-1-1-4-7-6", "CM-1-1-4-7-7", "CM-1-1-4-7-8",
              "CM-1-1-7-1-1", "CM-1-1-7-1-2", "CM-1-1-7-1-3",
              "CM-1-1-7-2-1", "CM-1-1-7-2-2", "CM-1-1-7-2-3", "CM-1-1-7-2-4",
              "CM-1-1-7-2-5", "CM-1-1-7-2-6",
              "CM-1-1-7-3-1", "CM-1-1-7-3-2", "CM-1-1-7-3-3"]:
    parent = child.rsplit("-", 1)[0]
    expected[child] = (parent, 4)
expected["CM-1-1-4-6"] = ("CM-1-1", 3)
expected["CM-1-1-4-6-1"] = ("CM-1-1-4-6", 4)
check("42 migration rows in spec", len(expected) == 42, f"{len(expected)}")
bad = []
for s, (p, lvl) in expected.items():
    r = im_new.get(s)
    if not r or r.get("parent_slug") != p or r.get("level") != lvl:
        bad.append((s, r.get("parent_slug") if r else None,
                    r.get("level") if r else None))
check("42 rows: correct new parent+level", not bad, str(bad[:5]))

# ---- identity stability vs baseline ----
new_slugs = {slug(c) for c in C_new} - {slug(c) for c in C_base}
check("slug set = base + 1 new L2",
      new_slugs == {"L2-refinery-asset-reliability-and-turnaround-coordination"},
      str(new_slugs))
pl_changes = [slug(c) for c in C_base
              if slug(c) in {slug(x) for x in C_new}
              and pref(g_new, URIRef(PROC + slug(c))) != pref(g_base, c)]
check("exactly 2 prefLabel changes",
      sorted(pl_changes) == ["CM-1-1-4", "CM-1-1-7"], str(pl_changes))
al_new = len(list(g_new.subject_objects(SKOS.altLabel)))
check("altLabel count 317", al_new == 317, f"{al_new}")
oldnames = {"Refinery Planning", "Refinery Scheduling"}
kept = [str(o) for o in g_new.objects(None, SKOS.altLabel) if str(o) in oldnames]
check("old promoted names not kept as altLabel", not kept, str(kept))
not_base = [slug(c) for c in C_base
            if slug(c) in {slug(x) for x in C_new}
            and (list(g_base.objects(c, SKOS.notation))
                 != list(g_new.objects(URIRef(PROC + slug(c)), SKOS.notation)))]
check("notations stable", not not_base, str(not_base[:5]))

# ---- tombstone + holds ----
t = URIRef(PROC + "CM-1-1-4-6")
check("tombstone deprecated",
      (t, RDF.type, OWL.DeprecatedClass) in g_new
      or bool(list(g_new.objects(t, OWL.deprecated))))
check("tombstone broader CM-1-1", slug(broader(g_new, t)) == "CM-1-1")
r61 = im_new.get("CM-1-1-4-6-1", {})
check("CM-1-1-4-6-1 level 4 under tombstone",
      r61.get("parent_slug") == "CM-1-1-4-6" and r61.get("level") == 4)

# ---- new L2 ----
NL2 = "L2-refinery-asset-reliability-and-turnaround-coordination"
nl2 = URIRef(PROC + NL2)
check("new L2 level 2 under L1-refining",
      slug(broader(g_new, nl2)) == "L1-refining"
      and im_new[NL2]["level"] == 2)
defs = list(g_new.objects(nl2, SKOS.definition))
scopes = list(g_new.objects(nl2, SKOS.scopeNote))
check("new L2 has definition + scope note", bool(defs) and bool(scopes))
check("new L2 unpopulated", not list(g_new.subjects(SKOS.broader, nl2)))
check("new L2 no notation", not list(g_new.objects(nl2, SKOS.notation)))

# ---- subjects changed outside migration set ----
# Bnode labels are serializer-run artifacts, so per-subject signatures are
# computed over bnode-canonicalized shapes (recursive (predicate, object)
# structure), not raw bnode ids.
mig = set(expected) | {NL2}


def canon_shape(g, node, _seen=None):
    from rdflib import BNode as _BNode, Literal as _Literal
    if _seen is None:
        _seen = {}
    if isinstance(node, _BNode):
        if node in _seen:
            return _seen[node]
        _seen[node] = ("bnode-ref", len(_seen))  # cycle guard
        items = tuple(sorted(
            (str(p), canon_shape(g, o, _seen))
            for p, o in g.predicate_objects(node)))
        _seen[node] = ("bnode", items)
        return _seen[node]
    kind = "lit" if isinstance(node, _Literal) else "uri"
    return (kind, str(node))


def sig(g, c):
    return sorted((str(p), canon_shape(g, o)) for p, o in g.predicate_objects(c))


changed_outside = {slug(c) for c in C_base
                   if slug(c) in {slug(x) for x in C_new}
                   and slug(c) not in mig
                   and sig(g_base, c) != sig(g_new, URIRef(PROC + slug(c)))}
# Sanctioned outside-set changes (every one enumerated in
# build/r1/r1-cell-diff.md as "mechanical stale-label correction (R1 §4a/4c)"):
# - CM-1-1-2-9-1: the reviewed R1 interface entry (Candidate package §1f).
# - 15 concepts: pure "Refinery Planning" -> "Refinery Planning and
#   Optimization" (and scheduling-label) substitutions in
#   definition/scopeNote text. No structural, identity, or substance change.
SANCTIONED_OUTSIDE = {
    "CM-1-1-2-9-1",
    "CM-1-1-1-1-4", "CM-1-1-2-10", "CM-1-1-2-12", "CM-1-1-3-7-2",
    "CM-1-2-1-1", "CM-1-2-1-1-1", "CM-1-2-1-1-4",
    "CM-1-2-5-1", "CM-1-2-5-1-2", "CM-1-2-5-1-3", "CM-1-2-5-1-4",
    "CM-1-2-5-2", "CM-1-2-5-2-1", "CM-1-2-5-2-3", "CM-1-2-5-4-2",
}
check("changed outside migration set == sanctioned set only",
      changed_outside == SANCTIONED_OUTSIDE, str(sorted(changed_outside)))
ALLOWED_OUTSIDE_PREDS = {
    str(SKOS.definition), str(SKOS.scopeNote), INTAKE + "relatedConcepts",
}
pred_violations = []
for s in changed_outside:
    so = dict(sig(g_base, URIRef(PROC + s)))
    sn = dict(sig(g_new, URIRef(PROC + s)))
    for p in set(so) | set(sn):
        if so.get(p) != sn.get(p) and p not in ALLOWED_OUTSIDE_PREDS:
            pred_violations.append((s, p))
check("outside-set changes touch only prose/interface predicates",
      not pred_violations, str(pred_violations[:5]))

# ---- reviewed interface materialization ----
RC = URIRef(INTAKE + "relatedConcepts")
for s, val in REVIEWED_PENDING_INTERFACES.items():
    got = [str(o) for o in g_new.objects(URIRef(PROC + s), RC)]
    check(f"reviewed interface emitted: {s}", got == [val], str(got))

# ---- Gate 2: path register ----
with open(R1 / "r1-path-compatibility-register.csv") as f:
    rd = list(csv.DictReader(f))
disp = [r for r in rd if "No consumer" in r.get("disposition", "")]
check("CSV 42 data rows", len(rd) == 42, f"{len(rd)}")
check("all rows no-consumer disposition", len(disp) == 42, f"{len(disp)}")

# ---- Gate 3: governance overlays ----
mig_ov = [s for s, r in im_new.items() if "r1_migration" in r]
check("43 r1_migration overlays (42 migration + 1 new-L2 creation)",
      len(mig_ov) == 43, f"{len(mig_ov)}")
pn = [s for s, r in im_new.items() if "prior_name" in r]
check("98 prior_name", len(pn) == 98, f"{len(pn)}")
ncn = [s for s, r in im_new.items() if "name_change_note" in r]
check("98 name_change_note", len(ncn) == 98, f"{len(ncn)}")
sha = [s for s, r in im_new.items() if "scoped_historical_alias" in r]
check("92 scoped_historical_alias", len(sha) == 92, f"{len(sha)}")
feed = [s for s in im_new if s.startswith("CM-1-2-5")]
feed_ok = all(im_new[s].get("parent_slug") == im_base[s].get("parent_slug")
              and im_new[s].get("level") == im_base[s].get("level")
              for s in feed if s in im_base)
check("feedstock cluster parent/level stable", feed_ok, f"{len(feed)} nodes")
irow = im_new.get("CM-1-1-2-9-1", {})
check("CM-1-1-2-9-1 placement stable",
      irow.get("parent_slug") == im_base["CM-1-1-2-9-1"].get("parent_slug"))

# ---- held-for-review invariant (workbook scan) ----
from openpyxl import load_workbook  # noqa: E402

wb = load_workbook(WORKBOOK, read_only=True, data_only=True)
ws = wb["Review & authoring"]
headers = [c.value for c in ws[1]]
idx = {h: i for i, h in enumerate(headers)}
pending_ifaces = {}
for r in ws.iter_rows(min_row=2, values_only=True):
    s = (r[idx["slug"]] or "").strip()
    st = (r[idx["status"]] or "").strip()
    rc = (r[idx["related_concepts"]] or "").strip()
    if s and rc and st not in ("approved", "blocked", "retired"):
        pending_ifaces[s] = rc
unreviewed = {s: v for s, v in pending_ifaces.items()
              if s not in REVIEWED_PENDING_INTERFACES}
check("no unreviewed pending-row interfaces (held for review)",
      not unreviewed, str(sorted(unreviewed)))
mismatch = {s: (pending_ifaces[s], REVIEWED_PENDING_INTERFACES[s])
            for s in pending_ifaces
            if s in REVIEWED_PENDING_INTERFACES
            and pending_ifaces[s] != REVIEWED_PENDING_INTERFACES[s]}
check("reviewed pending interfaces match reviewed values",
      not mismatch, str(mismatch))

# ---- Gate 4: report consistency ----
rep = (OUT / "step3-taxonomy-report.md").read_text(encoding="utf-8")
check("report: 683 concepts", "Concepts: 683" in rep)
check("report: 681 broader links", "681 broader links" in rep)
check("report: no stale 680 prose", "680 broader links" not in rep)
check("report: 14496 triples", "Triples: 14496" in rep)

fails = [r for r in results if not r[1]]
print(f"\n{len(results) - len(fails)}/{len(results)} checks PASS")
for name, ok, detail in results:
    print(("PASS " if ok else "FAIL ") + name
          + (f" [{detail}]" if detail and not ok else ""))
sys.exit(1 if fails else 0)
