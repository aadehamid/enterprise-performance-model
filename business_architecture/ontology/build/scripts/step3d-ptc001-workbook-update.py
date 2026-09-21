#!/usr/bin/env python3
"""Step 3d: PTC-001 workbook updates (partial resolution).

Hamid's decision (2026-09-21, revised proposal): the tree pass moved two
rows and tombstoned Commercial Development. This script mirrors those
changes into step3c-definition-authoring-workbook.xlsx:

  - CM-1-1-4-6-2 Plan Budgets: blocked -> pending, level 5 -> 3,
    new parent/breadcrumb, move note; draft definition pre-filled as a
    seed for the follow-up authoring batch.
  - CM-1-1-4-6-3 -> "Coordinate Site Business Risk Management":
    blocked -> pending, level 5 -> 3, renamed (prior name -> altLabel),
    new parent/breadcrumb, scope-boundary note; draft definition +
    scope note pre-filled as seeds.
  - CM-1-1-4-6-1 Develop Strategic Business Plan: stays blocked;
    terminology_notes cites PTC-001-B.
  - CM-1-1-4-6 Commercial Development: stays retired; note records the
    owl:deprecated tombstoning.
  - 2 new rows appended: the Candidate L2 architecture nodes
    (L2-financial-planning-and-performance-management,
    L2-refinery-performance-and-risk-coordination), status pending,
    draft definition + scope note pre-filled as seeds, Candidate note
    in terminology_notes.

Draft text on pending rows is fine and common (validator raises a
non-blocking question, which is intended here).

Breadcrumbs are recomputed from the new downstream_process_map.json
for all six affected rows.

One-shot. Asserts expected pre-state, so a stale workbook fails loudly.
Run from the repo root.
"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[4]
WB = ROOT / "business_architecture/ontology/build/output/step3c-definition-authoring-workbook.xlsx"
JSON = ROOT / "business_architecture" / "business_process" / "downstream_process_map.json"
SHEET = "Review & authoring"

FIN_L2 = "L2-financial-planning-and-performance-management"
FIN_L2_NAME = "Financial Planning and Performance Management"
REF_L2 = "L2-refinery-performance-and-risk-coordination"
REF_L2_NAME = "Refinery Performance and Risk Coordination"
RISK_NEW_NAME = "Coordinate Site Business Risk Management"
RISK_OLD_NAME = "Manage Site Specific Business Risk"

FIN_L2_DEF = (
    "The Finance-owned capability that establishes and operates the financial "
    "planning and performance-management framework; coordinates, consolidates, "
    "challenges, validates, publishes, monitors, and refreshes enterprise and "
    "business-unit financial plans and budgets in accordance with the approved "
    "planning calendar, planning assumptions, delegated authorities, and "
    "financial-control requirements; and measures and reports financial "
    "performance against those plans."
)
FIN_L2_SCOPE = (
    "Includes the financial planning calendar and planning assumptions; budget "
    "development, consolidation, and version governance; capital and operating "
    "expenditure planning controls; forecasting; financial performance reporting, "
    "variance analysis, escalation, and financial-plan review and "
    "approval-workflow administration. Operating domains own their underlying "
    "demand, volume, margin, capacity, operating-cost, capital, and resource "
    "assumptions and provide their plans and forecasts as inputs. Finance owns "
    "the financial planning process, consolidation, control framework, and "
    "financial reporting. Final approval remains with the applicable delegated "
    "authority. Excludes enterprise strategic direction, portfolio choices, and "
    "capital allocation decisions pending PTC-001-B. Excludes fixed-asset "
    "project accounting where intentionally governed and executed as an "
    "ERP-native capability."
)
REF_L2_DEF = (
    "The Refining-owned capability that coordinates refinery performance "
    "oversight and site-level business-risk management, including the "
    "identification, assessment, treatment coordination, monitoring, and "
    "escalation of risks that may affect refinery objectives."
)
REF_L2_SCOPE = (
    "Includes coordination of site business-risk ownership and treatment for "
    "risks affecting production reliability, throughput, feedstock and product "
    "quality, turnaround execution, cost, schedule, asset availability, supply, "
    "product disposition, and local stakeholder commitments. Does not operate "
    "refinery units; set enterprise risk appetite or policy; replace EHS or "
    "process-safety management; administer insurance; approve capital; manage "
    "financial-market risk; conduct Internal Audit; or own Legal, Compliance, "
    "Security, Cybersecurity, or business-continuity authorities."
)
BUDGET_DEF = (
    "Develop, coordinate, consolidate, challenge, approve, publish, monitor, "
    "and refresh enterprise and business-unit budgets in accordance with the "
    "approved financial-planning calendar, planning assumptions, delegated "
    "authorities, and financial-control requirements. Operating domains own "
    "their underlying demand, volume, margin, capacity, operating-cost, "
    "capital, and resource assumptions; Finance owns the financial planning "
    "process, consolidation, control, and reporting framework."
)
RISK_DEF = (
    "Identify, assess, document, prioritize, coordinate treatment of, monitor, "
    "and escalate site-specific business risks affecting refinery objectives, "
    "production reliability, cost, schedule, asset availability, supply, "
    "product disposition, and local stakeholder commitments, in accordance "
    "with enterprise risk policy, delegated authority, and specialist-control "
    "requirements."
)
RISK_SCOPE = (
    "Coordinates site-level risk ownership, treatment, monitoring, and "
    "escalation; it does not set enterprise risk appetite or policy, replace "
    "EHS / process-safety management, administer insurance, approve capital, "
    "manage financial-market risk, conduct Internal Audit, or own Legal, "
    "Compliance, Security, Cybersecurity, or business-continuity authorities."
)

CANDIDATE_NOTE = (
    "Step 3d PTC-001 tree pass (2026-09-21): new Candidate architecture node — "
    "structural placement approved via ptc-001-tree-proposal.md (Hamid). "
    "Definition authoring queued as a follow-up mini-batch."
)


def breadcrumbs(data):
    out = {}

    def walk(node, trail):
        trail = trail + (node["name"],)
        pid = node.get("id")
        slug = (pid.replace(" ", "-").replace(".", "-") if pid
                else f"L{node['level']}-" + node["name"].lower().replace(" ", "-").replace("&", "").strip("-"))
        # recompute stub slug exactly like step2
        if not pid:
            import re, unicodedata
            t = unicodedata.normalize("NFKD", node["name"]).encode("ascii", "ignore").decode().lower()
            t = re.sub(r"[^a-z0-9]+", "-", t).strip("-")
            slug = f"L{node['level']}-{t}"
        out[slug] = " > ".join(trail)
        for c in node.get("children", []) or []:
            walk(c, trail)

    for top in data:
        walk(top, ())
    return out


def main() -> int:
    data = json.loads(JSON.read_text(encoding="utf-8"))
    crumbs = breadcrumbs(data)

    wb = load_workbook(WB)
    ws = wb[SHEET]
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    by_slug = {}
    for r in ws.iter_rows(min_row=2):
        by_slug[r[idx["slug"]].value] = r

    # --- pre-state assertions ---
    assert len(by_slug) == 503, f"expected 503 rows, got {len(by_slug)}"
    r62 = by_slug["CM-1-1-4-6-2"]
    r63 = by_slug["CM-1-1-4-6-3"]
    r61 = by_slug["CM-1-1-4-6-1"]
    r6 = by_slug["CM-1-1-4-6"]
    assert r62[idx["status"]].value == "blocked" and r62[idx["level"]].value == 5
    assert r63[idx["status"]].value == "blocked" and r63[idx["name"]].value == RISK_OLD_NAME
    assert r61[idx["status"]].value == "blocked"
    assert r6[idx["status"]].value == "retired"
    assert FIN_L2 not in by_slug and REF_L2 not in by_slug

    def append_note(row, note):
        cur = row[idx["terminology_notes"]].value or ""
        row[idx["terminology_notes"]].value = (cur + " " + note).strip() if cur else note

    # --- Plan Budgets: blocked -> pending, moved ---
    r62[idx["status"]].value = "pending"
    r62[idx["level"]].value = 3
    r62[idx["parent"]].value = FIN_L2_NAME
    r62[idx["breadcrumb"]].value = crumbs["CM-1-1-4-6-2"]
    r62[idx["definition"]].value = BUDGET_DEF
    append_note(r62,
        "Step 3d PTC-001 tree pass (2026-09-21): reparented from Commercial "
        "Development (CM-1-1-4-6) to Financial Planning and Performance "
        "Management under Finance; level 5 -> 3. Slug/URI/notation unchanged. "
        "PTC-001 partially resolved.")

    # --- risk row: renamed, blocked -> pending, moved ---
    r63[idx["name"]].value = RISK_NEW_NAME
    r63[idx["alt_labels"]].value = RISK_OLD_NAME
    r63[idx["status"]].value = "pending"
    r63[idx["level"]].value = 3
    r63[idx["parent"]].value = REF_L2_NAME
    r63[idx["breadcrumb"]].value = crumbs["CM-1-1-4-6-3"]
    r63[idx["definition"]].value = RISK_DEF
    r63[idx["scope_note"]].value = RISK_SCOPE
    append_note(r63,
        "Step 3d PTC-001 tree pass (2026-09-21): reparented from Commercial "
        "Development to Refinery Performance and Risk Coordination under "
        "Refining; renamed from 'Manage Site Specific Business Risk' (kept as "
        "altLabel); level 5 -> 3. Scope boundary: coordinates site-level risk "
        "ownership/treatment/monitoring/escalation; does not set enterprise "
        "risk appetite or replace specialist authorities. PTC-001 partially "
        "resolved.")

    # --- strategy row: stays blocked, PTC-001-B note ---
    r61[idx["breadcrumb"]].value = crumbs["CM-1-1-4-6-1"]
    append_note(r61,
        "Step 3d PTC-001 tree pass (2026-09-21): stays blocked under PTC-001-B "
        "— enterprise strategy-ownership decision (options: Corporate Planning "
        "within Finance / Corporate Strategy / executive cross-functional "
        "governance). Parent Commercial Development tombstoned (owl:deprecated) "
        "pending PTC-001-B resolution.")

    # --- tombstone: stays retired, note records the tombstoning ---
    r6[idx["breadcrumb"]].value = crumbs["CM-1-1-4-6"]
    append_note(r6,
        "Step 3d PTC-001 tree pass (2026-09-21): tombstoned as owl:deprecated — "
        "removed from active navigation, retained for lineage (PTC-001 / "
        "PTC-001-B traceability). Physical removal is a later explicit "
        "retirement decision once PTC-001-B is resolved.")

    # --- new Candidate L2 rows ---
    def new_row(slug, name, parent_name, definition, scope):
        row = [None] * len(headers)
        row[idx["slug"]] = slug
        row[idx["level"]] = 2
        row[idx["name"]] = name
        row[idx["breadcrumb"]] = crumbs[slug]
        row[idx["parent"]] = parent_name
        row[idx["is_stub"]] = "no"
        row[idx["definition"]] = definition
        row[idx["scope_note"]] = scope
        row[idx["terminology_notes"]] = CANDIDATE_NOTE
        row[idx["status"]] = "pending"
        return row

    ws.append(new_row(FIN_L2, FIN_L2_NAME, "Finance", FIN_L2_DEF, FIN_L2_SCOPE))
    ws.append(new_row(REF_L2, REF_L2_NAME, "Refining", REF_L2_DEF, REF_L2_SCOPE))

    wb.save(WB)

    # --- post-state assertions ---
    wb2 = load_workbook(WB, read_only=True)
    ws2 = wb2[SHEET]
    n = ws2.max_row - 1
    assert n == 505, f"expected 505 rows, got {n}"
    print(f"workbook: 503 -> 505 rows")
    print("  ~ CM-1-1-4-6-2 Plan Budgets: blocked -> pending, moved to Finance")
    print(f"  ~ CM-1-1-4-6-3: renamed to {RISK_NEW_NAME!r}, blocked -> pending, moved to Refining")
    print("  ~ CM-1-1-4-6-1: stays blocked (PTC-001-B); CM-1-1-4-6: stays retired (tombstoned)")
    print("  + 2 Candidate L2 rows appended")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
