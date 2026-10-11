#!/usr/bin/env python3
"""Fixture tests for scope-note-append.py."""
from __future__ import annotations

import csv
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

SCRIPT = Path(__file__).resolve().parent / "scope-note-append.py"
HEADERS = ["slug", "name", "scope_note", "status"]


def _fixture(rows: list[dict], edits: list[tuple[str, str]]) -> tuple[Path, Path]:
    tmp = Path(tempfile.mkdtemp())
    wb = Workbook()
    ws = wb.active
    ws.title = "Review & authoring"
    ws.append(HEADERS)
    for row in rows:
        ws.append([row.get(h, "") for h in HEADERS])
    wb.save(tmp / "wb.xlsx")
    with open(tmp / "review.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["row_id", "scope_note_slug", "scope_note_sentence"])
        for slug, sentence in edits:
            w.writerow(["REL-00001", slug, sentence])
    return tmp / "wb.xlsx", tmp / "review.csv"


def _run(wb: Path, review: Path, *extra: str) -> subprocess.CompletedProcess:
    return subprocess.run(
        [sys.executable, str(SCRIPT), "--workbook", str(wb), "--corrections", str(review), *extra],
        capture_output=True, text=True,
    )


def _notes(wb: Path) -> dict:
    ws = load_workbook(wb)["Review & authoring"]
    return {r[0]: r[2] for r in ws.iter_rows(min_row=2, values_only=True)}


class ScopeNoteAppend(unittest.TestCase):
    def test_appends_to_approved_row_and_skips_blank_slug(self):
        wb, review = _fixture(
            [{"slug": "A-1", "scope_note": "Existing note.", "status": "approved"}],
            [("A-1", "New  sentence."), ("", "")],
        )
        self.assertEqual(_run(wb, review).returncode, 0)
        self.assertEqual(_notes(wb)["A-1"], "Existing note. New sentence.")

    def test_rerun_changes_nothing(self):
        wb, review = _fixture(
            [{"slug": "A-1", "scope_note": "Existing note.", "status": "approved"}],
            [("A-1", "New sentence.")],
        )
        _run(wb, review)
        before = wb.read_bytes()
        result = _run(wb, review)
        self.assertIn("0 added, 1 already present", result.stdout)
        self.assertEqual(wb.read_bytes(), before)

    def test_sentence_inside_a_longer_sentence_is_still_added(self):
        wb, review = _fixture(
            [{"slug": "A-1", "scope_note": "Desk B Settles trades.", "status": "approved"}],
            [("A-1", "Settles trades.")],
        )
        self.assertEqual(_run(wb, review).returncode, 0)
        self.assertEqual(_notes(wb)["A-1"], "Desk B Settles trades. Settles trades.")

    def test_whole_sentence_mid_note_counts_as_present(self):
        wb, review = _fixture(
            [{"slug": "A-1", "scope_note": "First. New sentence. Last.", "status": "approved"}],
            [("A-1", "New sentence.")],
        )
        self.assertIn("0 added, 1 already present", _run(wb, review).stdout)

    def test_refuses_row_not_approved_and_writes_nothing(self):
        wb, review = _fixture(
            [{"slug": "A-1", "scope_note": "Note.", "status": "approved"},
             {"slug": "B-1", "scope_note": "Note.", "status": "blocked"}],
            [("A-1", "Sentence."), ("B-1", "Sentence.")],
        )
        result = _run(wb, review)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("B-1: status is 'blocked'", result.stderr)
        self.assertEqual(_notes(wb)["A-1"], "Note.")

    def test_refuses_unknown_slug(self):
        wb, review = _fixture(
            [{"slug": "A-1", "scope_note": "Note.", "status": "approved"}],
            [("Z-9", "Sentence.")],
        )
        result = _run(wb, review)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("Z-9: not in the workbook", result.stderr)

    def test_check_writes_nothing(self):
        wb, review = _fixture(
            [{"slug": "A-1", "scope_note": "Note.", "status": "approved"}],
            [("A-1", "Sentence.")],
        )
        before = wb.read_bytes()
        self.assertEqual(_run(wb, review, "--check").returncode, 0)
        self.assertEqual(wb.read_bytes(), before)


if __name__ == "__main__":
    unittest.main()
