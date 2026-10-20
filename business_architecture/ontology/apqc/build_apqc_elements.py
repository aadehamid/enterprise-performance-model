"""Build the APQC element list (Step 7; EPM-DEC-001-0060).

    python3 build_apqc_elements.py            # read the workbook and print the figures
    python3 build_apqc_elements.py --write    # also write apqc-elements.ttl
    python3 build_apqc_elements.py --verify   # fail unless apqc-elements.ttl equals a fresh build

Reads the "Combined" sheet of the vendored APQC Downstream Petroleum PCF v7.2.2 workbook
(business_architecture/reference/) and writes one resource per element, described the way the
core release describes the elements it references (identifier, part of the PCF, title), plus
the element's name and hierarchy ID. The build fails unless every PCF ID is unique and every
element core 2.1.0 references is in the workbook. APQC material is used with attribution per
the APQC/IBM licence (business_architecture/reference/README.md).
"""
import argparse
import os
import sys
from collections import Counter

import openpyxl
from rdflib import Graph, Literal, Namespace, URIRef
from rdflib.namespace import DCTERMS, RDF, RDFS, SKOS, XSD

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, "..", "..", ".."))
WORKBOOK = os.path.join(REPO, "business_architecture", "reference", "APQC-PCF-Downstream-Petroleum-v7.2.2.xlsx")
CORE = os.path.join(HERE, "..", "releases", "core", "2.1.0", "core.ttl")
OUTPUT = os.path.join(HERE, "apqc-elements.ttl")
SRC = Namespace("https://w3id.org/lsc/ontology/source/")
PCF = SRC["apqc-pcf-7.2.2"]
# The notice and attribution the APQC/IBM licence requires on every copy (reference/README.md).
ATTRIBUTION = ("This work is based in part on an Industry Process Classification Framework (\"PCF\"), an open "
               "standard developed by APQC and IBM to promote benchmarking and best practices worldwide. The PCF "
               "is intended to facilitate organizational improvement through process management and benchmarking, "
               "regardless of industry, size, or geography. To download the full PCF or other industry-specific "
               "versions of the PCF, as well as associated measures and benchmarking, please visit www.apqc.org/pcf. "
               "©2025 APQC and IBM. The PCF is licensed perpetually, worldwide, royalty-free for use, copying, "
               "publication, modification, and derivative works, provided this notice and the attribution above "
               "travel with all copies.")
VERSION = "7.2.2"
ELEMENT = "https://w3id.org/lsc/ontology/source/apqc-pcf-7.2.2/element/"


def elements():
    wb = openpyxl.load_workbook(WORKBOOK, read_only=True)
    intro = " ".join(str(c) for r in wb["Introduction"].iter_rows(values_only=True) for c in r if c)
    if f"VERSION NUMBER {VERSION}" not in intro.upper().replace("  ", " "):
        sys.exit(f"the workbook does not say it is version {VERSION}")
    rows = list(wb["Combined"].iter_rows(values_only=True))
    header = rows[0]
    assert header[:3] == ("PCF ID", "Hierarchy ID", "Name"), header
    return [(int(r[0]), str(r[1]), str(r[2]).strip()) for r in rows[1:] if r[0] is not None]


def load():
    rows = elements()
    problems = [f"PCF ID {i} appears {n} times" for i, n in Counter(i for i, _, _ in rows).items() if n > 1]
    core = Graph().parse(CORE)
    referenced = {str(o) for o in core.objects(None, DCTERMS.references) if str(o).startswith(ELEMENT)}
    ids = {ELEMENT + str(i) for i, _, _ in rows}
    problems += [f"core references {e}, which is not in the workbook" for e in sorted(referenced - ids)]
    if problems:
        sys.exit("APQC elements do not reconcile:\n  " + "\n  ".join(problems))
    return rows, referenced


def build(rows):
    g = Graph()
    g.bind("src", SRC)
    for pcf_id, hierarchy, name in rows:
        e = URIRef(ELEMENT + str(pcf_id))
        g.add((e, RDF.type, DCTERMS.BibliographicResource))
        g.add((e, DCTERMS.identifier, Literal(str(pcf_id))))
        g.add((e, DCTERMS.isPartOf, PCF))
        g.add((e, DCTERMS.title, Literal(f"APQC Downstream Petroleum PCF v7.2.2, element {pcf_id}", lang="en")))
        g.add((e, RDFS.label, Literal(name, lang="en")))
        g.add((e, SKOS.notation, Literal(hierarchy, datatype=XSD.string)))
    return g


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--write", action="store_true")
    ap.add_argument("--verify", action="store_true")
    args = ap.parse_args()
    rows, referenced = load()
    text = ("# APQC Downstream Petroleum PCF v7.2.2 elements (Step 7; EPM-DEC-001-0060). Built by build_apqc_elements.py\n"
            "# from the vendored workbook. Do not edit by hand.\n#\n"
            + "".join(f"# {line}\n" for line in ATTRIBUTION.replace(". ", ".\n").split("\n")) + "#\n"
            + build(rows).serialize(format="turtle"))
    print(f"elements {len(rows)}; referenced by core 2.1.0 {len(referenced)}; not referenced {len(rows) - len(referenced)}")
    if args.verify:
        if not (os.path.exists(OUTPUT) and open(OUTPUT).read() == text):
            sys.exit("apqc-elements.ttl is not current: run build_apqc_elements.py --write")
        print("apqc-elements.ttl is current")
    if args.write:
        with open(OUTPUT, "w") as f:
            f.write(text)
        print(f"wrote {OUTPUT}")


if __name__ == "__main__":
    main()
